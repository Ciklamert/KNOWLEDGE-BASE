"""Turn EVERY document into Markdown before it goes to LightRAG (``to_markdown``), following the idea of Microsoft's
MarkItDown: one entry point, a registry of small converters, each deciding by file extension AND by the content
itself (magic numbers, zip contents, text shape), tried in order with a fallback to the next converter and finally
to plain text, and an output that keeps the document's structure (headings, lists, tables, links, code) while
spending as few tokens as possible (no page furniture, no repeated page headers and footers, no blank-line runs).

Formats: PDF, Word, PowerPoint, Excel, HTML, CSV/TSV, JSON, XML, RSS/Atom feeds, Jupyter notebooks, EPUB e-books,
e-mails (.eml, with their attachments), zip archives (every file inside), images (with their metadata), source code,
configuration and text files. A file with a wrong or missing extension is recognised by its content.

Differences from MarkItDown, on purpose: tables stay HTML, because LightRAG's native engine analyses HTML tables as
tables; pictures are kept as image files in a ``.textpack`` (Markdown + images) for the vision model instead of being
captioned during conversion; everything is model-free and uses only packages already bundled (pdfplumber, pypdfium2,
python-docx, python-pptx, openpyxl, Pillow and the standard library).

The converted Markdown is also kept in data/converted/ and can be previewed in the UI.
"""
import base64
import hashlib
import html
import io
import json
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

MIN_IMG_PX = 64          # LightRAG skips smaller images anyway (VLM_MIN_IMAGE_PIXEL)
MAX_SHEET_ROWS = 2000
MAX_SHEET_COLS = 40


class Pack:
    """Collects markdown + assets, writes a .textpack zip."""

    def __init__(self, title: str):
        self.md = [f"# {title}", ""]
        self.assets: dict[str, bytes] = {}
        self._seen: Counter = Counter()

    def add_image(self, data: bytes, ext: str, alt: str) -> None:
        h = hashlib.sha1(data).hexdigest()[:12]
        self._seen[h] += 1
        if self._seen[h] > 2:  # logos / headers repeated on every page: describe once or twice
            return
        name = f"assets/{h}.{ext}"
        self.assets[name] = data
        self.md += [f"![{alt}]({name})", ""]

    def build(self) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("text.md", "\n".join(self.md))
            for name, data in self.assets.items():
                z.writestr(name, data)
        return buf.getvalue()


def html_table(rows: list[list]) -> str:
    rows = [[("" if c is None else " ".join(str(c).split())) for c in r] for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    head = "".join(f"<th>{html.escape(c)}</th>" for c in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in rows[1:])
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _png(img) -> bytes:
    out = io.BytesIO()
    img.convert("RGB").save(out, "PNG", optimize=True)
    return out.getvalue()


def _jpg(img) -> bytes:
    out = io.BytesIO()
    img.convert("RGB").save(out, "JPEG", quality=82)
    return out.getvalue()


# ---------------------------------------------------------------- PDF
def pdf(data: bytes, title: str) -> "Pack":
    import pdfplumber
    import pypdfium2

    pack = Pack(title)
    doc = pypdfium2.PdfDocument(data)
    with pdfplumber.open(io.BytesIO(data)) as pdfp:
        sizes = Counter()
        for page in pdfp.pages[:30]:
            sizes.update(round(c["size"]) for c in page.chars[:4000])
        body = sizes.most_common(1)[0][0] if sizes else 10
        heading_sizes = sorted({s for s in sizes if s >= body * 1.18}, reverse=True)[:3]

        edges: list[set[str]] = []   # normalised first and last lines of each page (headers / footers)
        for pno, page in enumerate(pdfp.pages):
            scale = 2.0
            rendered = None

            def render():
                nonlocal rendered
                if rendered is None:
                    rendered = doc[pno].render(scale=scale).to_pil()
                return rendered

            text_chars = [c for c in page.chars if c.get("text", "").strip()]
            if len(text_chars) < 20:  # scanned page / pure figure: let the vision model read it
                pack.md += [f"## Page {pno + 1}", ""]
                pack.add_image(_jpg(render()), "jpg", f"Page {pno + 1} of {title} (scanned or graphic page)")
                continue

            blocks = []  # (top, kind, payload)
            tables = page.find_tables()
            tboxes = [t.bbox for t in tables]
            for t in tables:
                h = html_table(t.extract())
                if h:
                    blocks.append((t.bbox[1], "table", h))

            for im in page.images:
                x0, top, x1, bottom = im["x0"], im["top"], im["x1"], im["bottom"]
                if (x1 - x0) * scale < MIN_IMG_PX or (bottom - top) * scale < MIN_IMG_PX:
                    continue
                box = (max(0, x0) * scale, max(0, top) * scale, min(page.width, x1) * scale,
                       min(page.height, bottom) * scale)
                if box[2] - box[0] < MIN_IMG_PX or box[3] - box[1] < MIN_IMG_PX:
                    continue
                blocks.append((top, "image", _png(render().crop(tuple(int(v) for v in box)))))

            def inside(obj):
                cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
                return any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in tboxes)

            for line in page.extract_text_lines(return_chars=True, strip=True):
                if not line["text"].strip() or inside(line):
                    continue
                size = round(max((c["size"] for c in line["chars"]), default=body))
                blocks.append((line["top"], "line", (line["text"], size, line["bottom"])))

            # many vector strokes outside tables = chart/diagram drawn with paths -> page snapshot
            strokes = [o for o in page.curves + page.lines + page.rects if not inside(o)]
            if len(strokes) > 40:
                blocks.append((page.height, "image", _jpg(render())))

            blocks.sort(key=lambda b: b[0])
            texts = [p[0] for _, k, p in blocks if k == "line"]
            if texts:
                edges.append({_norm_line(t) for t in texts[:2] + texts[-2:]})
            para, last_bottom = [], None

            def flush():
                if para:
                    pack.md += [" ".join(para), ""]
                    para.clear()
            for top, kind, payload in blocks:
                if kind == "line":
                    text, size, bottom = payload
                    if size in heading_sizes and len(text) < 120:
                        flush()
                        pack.md += ["#" * (heading_sizes.index(size) + 2) + " " + text, ""]
                    else:
                        if last_bottom is not None and top - last_bottom > size * 0.9:
                            flush()
                        para.append(text)
                    last_bottom = bottom
                else:
                    flush()
                    last_bottom = None
                    if kind == "table":
                        pack.md += [payload, ""]
                    else:
                        pack.add_image(payload, "png" if payload[:4] == b"\x89PNG" else "jpg",
                                       f"Figure on page {pno + 1} of {title}")
            flush()
    doc.close()
    pack.md = _drop_repeated(pack.md, edges)
    return pack


def _norm_line(text: str) -> str:
    return re.sub(r"\d+", "#", " ".join(text.split()).lower())


def _drop_repeated(md: list[str], edges: list[set[str]]) -> list[str]:
    """Remove page headers and footers: lines at the top or bottom of at least 3 pages and half of all pages
    (page numbers count as the same line). They carry no meaning and cost tokens on every page."""
    counts = Counter(t for page in edges for t in page)
    repeated = {t for t, n in counts.items() if n >= 3 and n >= len(edges) / 2 and len(t) < 120}
    return [line for line in md if _norm_line(line) not in repeated] if repeated else md


# ---------------------------------------------------------------- PowerPoint
def pptx(data: bytes, title: str) -> "Pack":
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pack = Pack(title)
    prs = Presentation(io.BytesIO(data))

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh
    for n, slide in enumerate(prs.slides, 1):
        head = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        pack.md += [f"## Slide {n}{': ' + head if head else ''}", ""]
        for sh in sorted(walk(slide.shapes), key=lambda s: (s.top or 0, s.left or 0)):
            if slide.shapes.title is not None and sh.shape_id == slide.shapes.title.shape_id:
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        pack.md.append(("  " * p.level) + "- " + t)
                pack.md.append("")
            elif getattr(sh, "has_table", False) and sh.has_table:
                pack.md += [html_table([[c.text for c in r.cells] for r in sh.table.rows]), ""]
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                try:
                    cats = list(ch.plots[0].categories)
                    rows = [["series", *map(str, cats)]] + [[s.name, *s.values] for s in ch.plots[0].series]
                    pack.md += [f"Chart: {ch.chart_title.text_frame.text if ch.has_title else 'untitled'}", "",
                                html_table(rows), ""]
                except Exception:
                    pass
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pack.add_image(sh.image.blob, sh.image.ext.replace("jpeg", "jpg"), f"Image on slide {n}")
                except Exception:
                    pass
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            pack.md += ["Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip(), ""]
    return pack


# ---------------------------------------------------------------- Excel
def xlsx(data: bytes, title: str) -> "Pack":
    from openpyxl import load_workbook

    pack = Pack(title)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                break
            rows.append(list(r[:MAX_SHEET_COLS]))
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if not rows:
            continue
        pack.md += [f"## Sheet: {ws.title}", ""]
        # split very long sheets into tables of 200 rows (header repeated) so each stays analysable
        header, rest = rows[0], rows[1:] or []
        for i in range(0, max(len(rest), 1), 200):
            pack.md += [html_table([header] + rest[i:i + 200]), ""]
        if ws.max_row and ws.max_row > MAX_SHEET_ROWS:
            pack.md += [f"(sheet truncated: first {MAX_SHEET_ROWS} of {ws.max_row} rows)", ""]
    return pack


# ---------------------------------------------------------------- Word
def docx(data: bytes, title: str) -> "Pack":
    """Headings (from Word styles), lists, bold/italic, tables (HTML) and pictures, in document order."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    pack = Pack(title)
    doc = Document(io.BytesIO(data))
    rels = doc.part.related_parts

    def runs_md(p) -> str:
        out = []
        for r in p.runs:
            t = r.text
            if not t:
                continue
            if r.bold and t.strip():
                t = f"**{t.strip()}** "
            elif r.italic and t.strip():
                t = f"*{t.strip()}* "
            out.append(t)
        return "".join(out).strip()
    def end_list():
        if pack.md and pack.md[-1].startswith("- "):
            pack.md.append("")
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(block, doc)
            for blip in block.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                part = rels.get(rid)
                if part is not None:
                    ext = (part.partname.ext or "png").lower().replace("jpeg", "jpg")
                    if ext in ("png", "jpg", "gif", "webp"):
                        pack.add_image(part.blob, ext, f"Image in {title}")
            text = runs_md(p)
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.match(r"(heading|başlık)\s*(\d)", style)
            is_list = "list" in style or (p._p.pPr is not None and p._p.pPr.numPr is not None)
            if not is_list:
                end_list()
            if style == "title":
                pack.md += [f"# {text}", ""]
            elif m:
                pack.md += ["#" * min(6, int(m.group(2)) + 1) + " " + text, ""]
            elif is_list:
                pack.md.append(f"- {text}")
            else:
                pack.md += [text, ""]
        elif tag == "tbl":
            end_list()
            t = Table(block, doc)
            pack.md += [html_table([[c.text for c in row.cells] for row in t.rows]), ""]
    return pack


# ---------------------------------------------------------------- HTML
class _Html2Md:
    """Small, dependency-free HTML -> Markdown: headings, paragraphs, lists, links, emphasis, code blocks, image alt
    text and tables (kept as HTML). Page furniture (scripts, navigation, side bars, forms, footers) is dropped, as
    MarkItDown does, so only the content costs tokens."""
    SKIP = ("script", "style", "head", "nav", "noscript", "aside", "form", "footer", "svg", "button", "iframe", "select")

    def __init__(self):
        from html.parser import HTMLParser
        self.out: list[str] = []
        self.buf: list[str] = []
        self.table: list[list[str]] | None = None
        self.skip = 0
        self.pre = 0
        self.href: str | None = None
        self.list_depth = 0
        outer = self

        class P(HTMLParser):
            def handle_starttag(self, tag, attrs):
                outer.start(tag, dict(attrs))

            def handle_endtag(self, tag):
                outer.end(tag)

            def handle_data(self, data):
                outer.data(data)
        self.parser = P(convert_charrefs=True)

    def flush(self, prefix=""):
        text = " ".join("".join(self.buf).split())
        self.buf = []
        if text:
            self.out += [prefix + text] + ([] if prefix.lstrip().startswith("- ") else [""])   # list items stay together

    def start(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif self.skip:
            return
        elif tag == "pre":
            self.flush()
            self.pre += 1
            self.buf = []
        elif tag == "table":
            self.flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.table.append([])
        elif tag in ("td", "th") and self.table is not None:
            self.buf = []
        elif tag in ("ul", "ol"):
            self.flush("  " * max(0, self.list_depth - 1) + "- " if self.list_depth else "")   # text of the parent item
            self.list_depth += 1
        elif tag in ("p", "div", "br", "section", "article", "main", "header", "blockquote") or re.fullmatch(r"h[1-6]", tag) or tag == "li":
            self.flush()
        elif tag == "code" and not self.pre:
            self.buf.append("`")
        elif tag in ("b", "strong"):
            self.buf.append("**")
        elif tag in ("i", "em"):
            self.buf.append("*")
        elif tag == "a":
            href = attrs.get("href") or ""
            self.href = href if href.startswith(("http://", "https://", "mailto:")) else None
            if self.href:
                self.buf.append("[")
        elif tag == "img":
            alt = (attrs.get("alt") or "").strip()
            if alt:
                self.buf.append(f"(image: {alt})")

    def end(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif self.skip:
            return
        elif tag == "pre" and self.pre:
            self.pre -= 1
            code = "".join(self.buf).strip("\n")
            self.buf = []
            if code.strip():
                self.out += ["```", code, "```", ""]
        elif tag in ("td", "th") and self.table is not None:
            if self.table:
                self.table[-1].append(" ".join("".join(self.buf).split()))
            self.buf = []
        elif tag == "table" and self.table is not None:
            self.out += [html_table(self.table), ""]
            self.table = None
        elif re.fullmatch(r"h[1-6]", tag):
            self.flush("#" * int(tag[1]) + " ")
        elif tag == "li":
            self.flush("  " * max(0, self.list_depth - 1) + "- ")
        elif tag in ("ul", "ol"):
            self.flush()
            self.list_depth = max(0, self.list_depth - 1)
            if not self.list_depth:
                self.out.append("")
        elif tag in ("p", "div", "section", "article", "main", "header", "blockquote"):
            self.flush()
        elif tag == "code" and not self.pre:
            self.buf.append("`")
        elif tag in ("b", "strong"):
            self.buf.append("**")
        elif tag in ("i", "em"):
            self.buf.append("*")
        elif tag == "a" and self.href:
            self.buf.append(f"]({self.href})")
            self.href = None

    def data(self, d):
        if not self.skip:
            self.buf.append(d)

    def convert(self, html_text: str) -> str:
        self.parser.feed(html_text)
        self.flush()
        md = "\n".join(self.out)
        return re.sub(r"\[\]\([^)]*\)", "", md)   # links without text (icons) left by the markup


def html_md(data: bytes, title: str) -> "Pack":
    pack = Pack(title)
    pack.md += [_Html2Md().convert(_decode(data))]
    return pack


# ---------------------------------------------------------------- plain text, tables, code, data
CODE_EXT = {".py": "python", ".java": "java", ".js": "javascript", ".ts": "typescript", ".cs": "csharp", ".go": "go",
            ".rb": "ruby", ".php": "php", ".kt": "kotlin", ".swift": "swift", ".c": "c", ".h": "c", ".cpp": "cpp",
            ".hpp": "cpp", ".rs": "rust", ".scala": "scala", ".sql": "sql", ".sh": "bash", ".ps1": "powershell",
            ".yaml": "yaml", ".yml": "yaml", ".toml": "toml", ".ini": "ini", ".properties": "properties",
            ".conf": "", ".cfg": "", ".gradle": "groovy", ".vb": "vbnet"}
TEXT_EXT = {".txt", ".text", ".log", ".rst", ".adoc", ".asciidoc", ".org", ".tex"}
MAX_TEXT = 2_000_000   # characters; longer plain text is cut (the cut is stated in the output)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1254", "latin-1"):   # UTF-8 first, then Turkish Windows, then anything
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def text_md(data: bytes, title: str) -> "Pack":
    pack = Pack(title)
    text = _decode(data).strip()
    pack.md += [text[:MAX_TEXT], ""] + ([f"(text cut at {MAX_TEXT:,} characters)", ""] if len(text) > MAX_TEXT else [])
    return pack


def csv_md(data: bytes, title: str, delimiter: str) -> "Pack":
    import csv
    pack = Pack(title)
    rows = list(csv.reader(io.StringIO(_decode(data)), delimiter=delimiter))[:MAX_SHEET_ROWS]
    if rows:
        header, rest = rows[0], rows[1:]
        for i in range(0, max(len(rest), 1), 200):
            pack.md += [html_table([header] + rest[i:i + 200]), ""]
    return pack


def code_md(data: bytes, title: str, lang: str) -> "Pack":
    pack = Pack(title)
    pack.md += [f"```{lang}", _decode(data).rstrip()[:MAX_TEXT], "```", ""]
    return pack


def json_md(data: bytes, title: str) -> "Pack":
    """JSON, pretty-printed in a fenced block (invalid JSON is kept as it is)."""
    text = _decode(data)
    try:
        text = json.dumps(json.loads(text), indent=1, ensure_ascii=False)
    except ValueError:
        pass
    pack = Pack(title)
    pack.md += ["```json", text[:MAX_TEXT], "```", ""]
    return pack


def notebook_md(data: bytes, title: str) -> "Pack":
    """Jupyter notebook: Markdown cells as they are, code cells fenced with the notebook's language, text outputs
    (shortened) and image outputs."""
    nb = json.loads(_decode(data))
    lang = ((nb.get("metadata") or {}).get("kernelspec") or {}).get("language") \
        or ((nb.get("metadata") or {}).get("language_info") or {}).get("name") or ""
    pack = Pack(title)
    src = lambda x: "".join(x) if isinstance(x, list) else str(x or "")
    for cell in nb.get("cells", []):
        if cell.get("cell_type") == "markdown":
            pack.md += [src(cell.get("source")).strip(), ""]
        elif cell.get("cell_type") == "code" and src(cell.get("source")).strip():
            pack.md += [f"```{lang}", src(cell.get("source")).rstrip(), "```", ""]
            for out in cell.get("outputs", []):
                d = out.get("data") or {}
                text = src(out.get("text")) or src(d.get("text/plain"))
                if text.strip():
                    lines = text.strip().splitlines()
                    pack.md += ["Output:", "```", "\n".join(lines[:20]) + ("\n…" if len(lines) > 20 else ""), "```", ""]
                if d.get("image/png"):
                    pack.add_image(base64.b64decode(src(d["image/png"])), "png", f"Chart or figure in {title}")
    return pack


def feed_md(data: bytes, title: str) -> "Pack":
    """RSS or Atom feed: one section per item with its date, link and summary."""
    import xml.etree.ElementTree as ET
    root = ET.fromstring(data)
    strip = lambda tag: tag.rsplit("}", 1)[-1]
    find = lambda el, name: next((c for c in el if strip(c.tag) == name), None)
    text = lambda el, name: ((find(el, name).text if find(el, name) is not None else "") or "").strip()
    channel = find(root, "channel") if strip(root.tag) == "rss" else root
    pack = Pack(text(channel, "title") or title)
    for item in (c for c in channel if strip(c.tag) in ("item", "entry")):
        link = text(item, "link") or ((find(item, "link").get("href") if find(item, "link") is not None else "") or "")
        summary = text(item, "description") or text(item, "summary") or text(item, "content")
        pack.md += [f"## {text(item, 'title') or 'Item'}", ""]
        meta = " · ".join(x for x in (text(item, "pubDate") or text(item, "updated") or text(item, "published"), link) if x)
        if meta:
            pack.md += [meta, ""]
        if summary:
            pack.md += [_Html2Md().convert(summary), ""]
    return pack


def xml_md(data: bytes, title: str) -> "Pack":
    pack = Pack(title)
    pack.md += ["```xml", _decode(data).strip()[:MAX_TEXT], "```", ""]
    return pack


def image_md(data: bytes, title: str, ext: str) -> "Pack":
    """The picture itself (for the vision model) plus its metadata as text: size, date, camera, description."""
    pack = Pack(title)
    try:
        from PIL import ExifTags, Image
        img = Image.open(io.BytesIO(data))
        exif = {ExifTags.TAGS.get(k, k): v for k, v in (img.getexif() or {}).items()}
        facts = [f"{img.width} × {img.height} pixels"] + [f"{k}: {exif[k]}" for k in (
            "ImageDescription", "DateTime", "Make", "Model", "Artist", "Copyright") if str(exif.get(k) or "").strip()]
        pack.md += ["- " + "\n- ".join(facts), ""]
    except Exception:
        pass
    pack.add_image(data, "jpg" if ext in (".jpg", ".jpeg") else ext.lstrip(".") or "png", title)
    return pack


# ---------------------------------------------------------------- containers: zip, EPUB, e-mail
MAX_MEMBERS = 200
MAX_MEMBER_BYTES = 25 * 1024 * 1024


def _merge(pack: "Pack", name: str, inner: "Pack") -> None:
    """Add a converted member under its own heading; its headings move one level down."""
    pack.md += [f"## {name}", ""] + [("#" + line) if line.startswith("#") else line for line in inner.md[1:]] + [""]
    pack.assets.update(inner.assets)


def zip_md(data: bytes, title: str, depth: int = 0) -> "Pack":
    """A zip archive: every supported file inside is converted, each under its own heading (one level deep)."""
    pack = Pack(title)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        members = [i for i in z.infolist() if not i.is_dir() and not i.filename.startswith(("__MACOSX/", "."))
                   and i.file_size <= MAX_MEMBER_BYTES][:MAX_MEMBERS]
        skipped = []
        for info in members:
            try:
                _merge(pack, info.filename, convert(info.filename, z.read(info), depth + 1))
            except Exception as exc:
                skipped.append(f"{info.filename} ({exc})"[:160])
        if skipped:
            pack.md += ["## Not converted", ""] + [f"- {s}" for s in skipped] + [""]
    return pack


def epub_md(data: bytes, title: str) -> "Pack":
    """An e-book: title and author from the package file, then the chapters in reading order."""
    import posixpath
    import xml.etree.ElementTree as ET
    z = zipfile.ZipFile(io.BytesIO(data))
    strip = lambda tag: tag.rsplit("}", 1)[-1]
    container = ET.fromstring(z.read("META-INF/container.xml"))
    opf_path = next(el.get("full-path") for el in container.iter() if strip(el.tag) == "rootfile")
    opf = ET.fromstring(z.read(opf_path))
    meta = {strip(el.tag): (el.text or "").strip() for el in opf.iter() if strip(el.tag) in ("title", "creator")}
    pack = Pack(meta.get("title") or title)
    if meta.get("creator"):
        pack.md += [f"Author: {meta['creator']}", ""]
    items = {el.get("id"): el.get("href") for el in opf.iter() if strip(el.tag) == "item"}
    base = posixpath.dirname(opf_path)
    for ref in (el.get("idref") for el in opf.iter() if strip(el.tag) == "itemref"):
        href = items.get(ref)
        if href:
            try:
                pack.md += [_Html2Md().convert(_decode(z.read(posixpath.join(base, href)))), ""]
            except KeyError:
                continue
    return pack


def email_md(data: bytes, title: str, depth: int = 0) -> "Pack":
    """An e-mail (.eml): sender, recipients, date and subject, the text (HTML converted), and every attachment that
    can be converted, each under its own heading."""
    from email import policy
    from email.parser import BytesParser
    msg = BytesParser(policy=policy.default).parsebytes(data)
    pack = Pack(str(msg.get("subject") or title))
    pack.md += [f"- {h}: {msg.get(h)}" for h in ("From", "To", "Cc", "Date") if msg.get(h)] + [""]
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        pack.md += [_Html2Md().convert(content) if body.get_content_type() == "text/html" else content.strip(), ""]
    for part in msg.iter_attachments():
        name = part.get_filename() or "attachment"
        try:
            _merge(pack, f"Attachment: {name}", convert(name, part.get_payload(decode=True) or b"", depth + 1))
        except Exception as exc:
            pack.md += [f"(attachment {name} not converted: {exc})", ""]
    return pack


# ---------------------------------------------------------------- the converter registry (the MarkItDown idea)
def sniff(data: bytes) -> str:
    """What the bytes are, whatever the file name says: magic numbers, zip contents, then text shape."""
    head = data[:4096]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            return "binary"
        if "word/document.xml" in names:
            return "docx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "META-INF/container.xml" in names:
            return "epub"
        return "zip"
    if head.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8")) or (head[:4] == b"RIFF" and head[8:12] == b"WEBP"):
        return "image"
    if b"\x00" in head:
        return "binary"
    text = head.decode("utf-8", "ignore").lstrip("﻿ \r\n\t").lower()
    if text.startswith("{") and '"cells"' in text and '"nbformat"' in _decode(data[:200000]):
        return "notebook"
    if text.startswith(("{", "[")):
        return "json"
    if text.startswith("<?xml") or text.startswith("<rss") or text.startswith("<feed"):
        return "feed" if re.search(r"<(rss|feed)[\s>]", text) else "xml"
    if text.startswith(("<!doctype html", "<html")) or "<body" in text:
        return "html"
    if re.match(r"(?:[a-z-]+: .*\r?\n)+", text) and re.search(r"^(from|subject|date): ", text, flags=re.M):
        return "eml"
    return "text"


@dataclass
class Converter:
    """One converter: the extensions and sniffed kinds it accepts, and its function (data, title, ext, depth) -> Pack.
    Converters are tried in order: the ones the extension AND the content agree on, then the content, then the
    extension; if one fails the next is tried, and plain text is the last resort (as in MarkItDown)."""
    name: str
    exts: tuple
    kinds: tuple
    fn: Callable


CONVERTERS = [
    Converter("pdf", (".pdf",), ("pdf",), lambda d, t, e, n: pdf(d, t)),
    Converter("word", (".docx",), ("docx",), lambda d, t, e, n: docx(d, t)),
    Converter("powerpoint", (".pptx",), ("pptx",), lambda d, t, e, n: pptx(d, t)),
    Converter("excel", (".xlsx", ".xlsm"), ("xlsx",), lambda d, t, e, n: xlsx(d, t)),
    Converter("epub", (".epub",), ("epub",), lambda d, t, e, n: epub_md(d, t)),
    Converter("notebook", (".ipynb",), ("notebook",), lambda d, t, e, n: notebook_md(d, t)),
    Converter("e-mail", (".eml",), ("eml",), lambda d, t, e, n: email_md(d, t, n)),
    Converter("feed", (".rss", ".atom"), ("feed",), lambda d, t, e, n: feed_md(d, t)),
    Converter("html", (".html", ".htm", ".xhtml"), ("html",), lambda d, t, e, n: html_md(d, t)),
    Converter("image", (".png", ".jpg", ".jpeg", ".webp", ".gif"), ("image",), lambda d, t, e, n: image_md(d, t, e)),
    Converter("csv", (".csv", ".tsv"), (), lambda d, t, e, n: csv_md(d, t, "\t" if e == ".tsv" else ",")),
    Converter("json", (".json",), ("json",), lambda d, t, e, n: json_md(d, t)),
    Converter("xml", (".xml",), ("xml",), lambda d, t, e, n: xml_md(d, t)),
    Converter("code", tuple(CODE_EXT), (), lambda d, t, e, n: code_md(d, t, CODE_EXT[e])),
    Converter("zip", (".zip",), ("zip",), lambda d, t, e, n: zip_md(d, t, n)),
    Converter("text", tuple(TEXT_EXT), ("text",), lambda d, t, e, n: text_md(d, t)),
]
SUPPORTED = sorted({".md", ".markdown"} | {e for c in CONVERTERS for e in c.exts})


def _tidy(lines: list[str]) -> list[str]:
    """Token-efficient Markdown: no trailing spaces, no runs of blank lines."""
    out: list[str] = []
    for line in "\n".join(lines).splitlines():
        line = line.rstrip()
        if line or (out and out[-1]):
            out.append(line)
    return out


def convert(filename: str, data: bytes, depth: int = 0) -> "Pack":
    """Any supported document -> Pack (Markdown + images). Raises ValueError for content no converter accepts."""
    ext = Path(filename).suffix.lower()
    title = re.sub(r"[_\-]+", " ", Path(filename).stem).strip() or filename
    if ext in (".md", ".markdown"):
        pack = Pack(title)
        pack.md = _decode(data).splitlines()
        return pack
    if depth > 1 and ext in (".zip", ".eml"):
        raise ValueError("nested archive")
    kind = sniff(data)
    # order: extension and content agree > a specific content type (a .xml that is an RSS feed, a .pdf that is
    # really a Word file) > the extension > a generic content type (json, xml, text) as the last resort
    both = [c for c in CONVERTERS if ext in c.exts and kind in c.kinds]
    by_kind = [c for c in CONVERTERS if kind in c.kinds and c not in both]
    by_ext = [c for c in CONVERTERS if ext in c.exts and c not in both and kind != "binary"]
    generic = kind in ("json", "xml", "text")
    errors = []
    for c in both + ([] if generic else by_kind) + by_ext + (by_kind if generic else []):
        try:
            pack = c.fn(data, title, ext, depth)
            pack.md = _tidy(pack.md)
            return pack
        except Exception as exc:   # the next converter may still read it (as in MarkItDown)
            errors.append(f"{c.name}: {type(exc).__name__}: {str(exc)[:120]}")
    raise ValueError(f"{filename}: cannot convert ({kind} content; " + ("; ".join(errors) or "unsupported format")
                     + f"). Supported: {', '.join(SUPPORTED)}")


def to_markdown(filename: str, data: bytes) -> tuple[str, bytes, str]:
    """EVERY document becomes Markdown before it reaches LightRAG.
    Returns (stored file name, payload, markdown text): `<file>.md` when the document has no images, else
    `<file>.textpack` (Markdown + image files, read by LightRAG's native engine; images go to the vision model)."""
    ext = Path(filename).suffix.lower()
    pack = convert(filename, data)
    md = "\n".join(pack.md)
    if pack.assets:
        return f"{filename}.textpack", pack.build(), md
    if ext in (".md", ".markdown"):
        return (filename if ext == ".md" else f"{filename}.md"), md.encode("utf-8"), md
    return f"{filename}.md", md.encode("utf-8"), md
