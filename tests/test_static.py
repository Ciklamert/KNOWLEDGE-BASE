"""Static checks: no server, no network.   python tests/test_static.py   (exit code = failures)"""
import hashlib
import importlib
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
_TMP = Path(tempfile.mkdtemp())
os.environ["BRAIN_CONFIG_FILE"] = str(_TMP / "test.env")   # never touch the real .env or data/
os.environ["BRAIN_DATA_DIR"] = str(_TMP / "data")
os.environ["BRAIN_REPOS_FILE"] = str(_TMP / "repos.json")
FAILS = []


def check(name, cond, info=""):
    print(f"{'PASS' if cond else 'FAIL'}  {name}  {info}")
    if not cond:
        FAILS.append(name)


def _raises(fn) -> bool:
    try:
        fn()
    except Exception:
        return True
    return False


def main():
    check(f"python 3.13 ({sys.version.split()[0]})", sys.version_info[:2] == (3, 13))
    for mod in ("brain.settings", "brain.query", "brain.jobs", "brain.llm", "brain.auth", "brain.search", "brain.repos", "brain.drift", "brain.agents",
                "brain.webhooks", "brain.mcp", "brain.server", "brain.codevec", "brain.srs", "brain.providers.codebase_memory", "brain.providers.docgraph",
                "brain.providers.convert", "brain.vcache", "brain.ocr"):
        try:
            importlib.import_module(mod)
            check(f"import {mod}", True)
        except Exception as exc:
            check(f"import {mod}", False, exc)

    from brain import providers, search
    from brain import settings as cfg
    check("code provider loads via registry", providers.code().name == "codebase_memory")
    check("doc provider loads via registry", providers.docs().name == "docgraph")

    sys.path.insert(0, str(ROOT / "tools"))
    import gen_config
    check(".env.example and 03-CONFIGURATION.md are up to date", gen_config.main(check=True) == 0)

    api_doc = (ROOT / "docs" / "06-API.md").read_text(encoding="utf-8")
    from brain.server import app
    for r in app.routes:
        path = getattr(r, "path", "")
        if path.startswith(("/api", "/webhooks", "/mcp", "/healthz")) and path != "/api/docs-ui":
            check(f"route documented: {path}", path.replace(":path}", "}") in api_doc)

    mcp_doc = (ROOT / "docs" / "07-MCP-AND-AGENTS.md").read_text(encoding="utf-8")
    from brain.mcp import TOOLS
    for t in TOOLS:
        check(f"MCP tool documented: {t['name']}", f"`{t['name']}`" in mcp_doc)

    from brain import query
    a = query.analyze("Müşteriden nasıl ödeme alınır ve KDV nerede uygulanır? chargeCustomer tax_rate")
    kw = a["keywords"]
    check("Turkish + identifier key terms", {"Müşteriden", "ödeme", "KDV", "chargeCustomer", "tax_rate"} <= set(kw), kw)
    check("identifiers ranked first", kw[:2] == ["chargeCustomer", "tax_rate"], kw)
    check("stop-words extracted", {"nasıl", "ve", "nerede"} <= set(a["stop_words"]) and "nasıl" not in kw, a["stop_words"])
    check("Turkish stem + identifier-part variants", {"Müşteri", "charge", "Customer", "tax", "rate"} <= set(a["variants"]), a["variants"])
    check("language detected", a["language"] == "tr", a["language"])
    b = query.analyze('How are "partial refunds" applied to invoices?')
    check("quoted phrase kept and ranked first", b["keywords"][0] == "partial refunds" and "invoice" in b["variants"], b)
    cfg.update({"BRAIN_STOPWORDS_EXTRA": "invoices"})
    check("extra stop-words setting", "invoices" in query.analyze("invoices total")["stop_words"])
    cfg.update({"BRAIN_STOPWORDS_EXTRA": ""})

    for key, bad in (("BRAIN_PORT", "abc"), ("BRAIN_VERIFY_TLS", "maybe"), ("CBM_INDEX_MODE", "turbo"), ("NOPE", "1")):
        try:
            cfg.validate(key, bad)
            check(f"rejects {key}={bad}", False)
        except ValueError:
            check(f"rejects {key}={bad}", True)
    scopes = cfg.update({"BRAIN_SEARCH_TOP_K": "5", "CBM__SOME_OPTION": "0"})
    check("update writes file and reports restarts", cfg.get("BRAIN_SEARCH_TOP_K") == 5 and "code" in scopes, scopes)
    check("pass-through variables", cfg.passthrough("CBM__").get("SOME_OPTION") == "0")
    check("unknown engine variables rejected", _raises(lambda: cfg.validate("OTHER__X", "1")))

    try:
        cfg.resolve_url("gitlab:a/b")
        check("gitlab: shorthand needs BRAIN_GITLAB_URL", False)
    except ValueError:
        check("gitlab: shorthand needs BRAIN_GITLAB_URL", True)
    cfg.update({"BRAIN_GITLAB_URL": "https://gitlab.corp.local/", "BRAIN_GITLAB_TOKEN": "gl", "BRAIN_TFS_URL":
                "https://tfs.corp.local/tfs/DefaultCollection", "BRAIN_TFS_TOKEN": "tf", "BRAIN_GIT_TOKEN": "gen"})
    check("gitlab: shorthand", cfg.resolve_url("gitlab:pay/billing") == "https://gitlab.corp.local/pay/billing.git")
    check("tfs: shorthand", cfg.resolve_url("tfs:Portal/web") == "https://tfs.corp.local/tfs/DefaultCollection/Portal/_git/web")
    check("full URL unchanged", cfg.resolve_url("https://x/y.git") == "https://x/y.git")
    check("GitLab credentials chosen by host", cfg.git_credentials("https://gitlab.corp.local/pay/billing.git")[1] == "gl")
    check("TFS credentials chosen by host", cfg.git_credentials("https://tfs.corp.local/tfs/DefaultCollection/P/_git/r")[1] == "tf")
    check("fallback git credentials", cfg.git_credentials("https://other/x.git")[1] == "gen")
    check("{gitlab: ...} field accepted", cfg.normalize_repo({"gitlab": "pay/billing"})["url"] == "gitlab:pay/billing"
          and cfg.normalize_repo({"gitlab": "pay/billing"})["name"] == "billing")
    from brain import webhooks
    import json as _json
    cfg.repos_file().write_text(_json.dumps({"repos": [{"name": "billing", "url": "gitlab:pay/billing"},
                                                       {"name": "web", "url": "tfs:Portal/web"}]}), encoding="utf-8")
    m1 = webhooks.match_repo(["https://gitlab.corp.local/pay/billing.git"])
    m2 = webhooks.match_repo(["https://tfs.corp.local/tfs/DefaultCollection/Portal/_git/web"])
    check("webhooks match shorthand repositories", (m1 or {}).get("name") == "billing" and (m2 or {}).get("name") == "web")

    # repositories, sets, per-repository webhooks
    saved = cfg.save_repos([{"name": "billing", "url": "gitlab:pay/billing"}, {"name": "web", "url": "tfs:Portal/web"}],
                           {"payments": ["billing", "web"]})
    sec = saved["repos"][0]["webhook_secret"]
    check("webhook secret generated per repository", len(sec) == 32 and sec != saved["repos"][1]["webhook_secret"])
    check("providers detected", [r["provider"] for r in saved["repos"]] == ["gitlab", "tfs"])
    check("sets stored", cfg.load_sets() == {"payments": ["billing", "web"]})
    again = cfg.save_repos([{"name": "billing", "url": "gitlab:pay/billing", "webhook_secret": ""}], {})
    check("webhook secret kept on re-save", again["repos"][0]["webhook_secret"] == sec)
    cfg.save_repos(saved["repos"], {"payments": ["billing", "web"]})
    code, body = webhooks.repo_hook("billing", {"x-gitlab-token": "wrong"}, {"object_kind": "push"})
    check("own webhook rejects a wrong GitLab secret", code == 401, body)
    import base64 as _b64
    auth = "Basic " + _b64.b64encode(f"x:{saved['repos'][1]['webhook_secret']}".encode()).decode()
    code, body = webhooks.repo_hook("web", {"authorization": auth}, {"eventType": "git.push", "resource": {"refUpdates": [{"name": "refs/heads/dev"}]}})
    check("own webhook accepts TFS secret (other branch ignored)", code == 200 and "ignored" in body, body)
    import time as _t
    from brain import repos as _repos
    gone = []
    real_rm, real_git = _repos.remove_branch, _repos.git
    _repos.remove_branch = lambda n, b: (gone.append(b), _repos._save_state(lambda d: d["branches"].pop(f"{n}@{b}", None)))

    def _overlay(repo, branch, age_days=0):
        _repos._save_state(lambda d: d["branches"].__setitem__(f"{repo}@{branch}", {
            "repo": repo, "branch": branch, "project": f"{repo}--{branch}", "path": str(_TMP / "wt"),
            "indexed_at": _t.time() - age_days * 86400}))
    _overlay("billing", "feature/x")
    code, body = webhooks.repo_hook("billing", {"x-gitlab-token": sec},
                                    {"object_kind": "push", "ref": "refs/heads/feature/x", "after": "0" * 40})
    check("GitLab branch deletion removes its overlay", code == 200 and gone == ["feature/x"], body)
    _overlay("web", "dev")
    code, body = webhooks.repo_hook("web", {"authorization": auth}, {"eventType": "git.push", "resource": {
        "refUpdates": [{"name": "refs/heads/dev", "newObjectId": "0" * 40}]}})
    check("TFS branch deletion removes its overlay", gone[-1:] == ["dev"] and "web@dev" not in _repos.state()["branches"], body)
    _overlay("billing", "kept")
    _overlay("billing", "deleted-quietly")
    _repos.git = lambda *a, **k: "1a2b\trefs/heads/main\n3c4d\trefs/heads/kept"
    check("reconcile: overlays whose branch is gone on the server are removed",
          _repos.reconcile_branches("billing") == ["deleted-quietly"] and "billing@kept" in _repos.state()["branches"])
    _overlay("billing", "old", age_days=40)
    cfg.update({"BRAIN_BRANCH_TTL_DAYS": "30"})
    check("expiry: an overlay unused for BRAIN_BRANCH_TTL_DAYS is removed", _repos.expire_branches() == ["billing@old"]
          and "billing@kept" in _repos.state()["branches"])
    _repos.remove_branch, _repos.git = real_rm, real_git
    _repos._save_state(lambda d: d["branches"].clear())

    # several main branches: configured per repository or by default, kept for ever, updated by webhooks,
    # and a feature branch is compared with the main it derives from
    check("other main branches: list or comma-separated, the main branch itself left out",
          cfg.normalize_repo({"name": "x", "path": "p", "branch": "main", "mains": "develop, main ,release/2.0"})["mains"]
          == ["develop", "release/2.0"])
    cfg.update({"BRAIN_OTHER_MAIN_BRANCHES": "develop"})
    check("other main branches: BRAIN_OTHER_MAIN_BRANCHES when a repository names none",
          cfg.normalize_repo({"name": "x", "path": "p"})["mains"] == ["develop"])
    cfg.update({"BRAIN_OTHER_MAIN_BRANCHES": ""})
    cfg.save_repos([{**saved["repos"][0], "mains": ["develop"]}, saved["repos"][1]], {"payments": ["billing", "web"]})
    _overlay("billing", "develop", age_days=400)
    _repos._save_state(lambda d: d["branches"]["billing@develop"].update(main=True))
    check("other main branches never expire", _repos.expire_branches() == [] and _repos.expires_at(
        _repos.state()["branches"]["billing@develop"]) is None)
    asked = []
    real_req = _repos.request_branch
    _repos.request_branch = lambda n, b, reason="", base=None: asked.append(b)
    code, body = webhooks.repo_hook("billing", {"x-gitlab-token": sec}, {"object_kind": "merge_request", "object_attributes": {
        "action": "merge", "state": "merged", "target_branch": "develop"}})
    check("a merge into another main branch updates it", asked == ["develop"] and code == 202, body)
    _repos.request_branch = real_req
    _repos._save_state(lambda d: d["branches"].clear())
    import subprocess as _sp
    origin, work = _TMP / "origin.git", _TMP / "work"
    g = lambda cwd, *a: _sp.run(["git", "-C", str(cwd), *a], capture_output=True, text=True, check=True).stdout.strip()
    _sp.run(["git", "init", "-q", "--bare", "-b", "main", str(origin)], check=True)
    _sp.run(["git", "clone", "-q", str(origin), str(work)], check=True, capture_output=True)
    for args in (["config", "user.email", "t@t"], ["config", "user.name", "t"], ["checkout", "-qb", "main"],
                 ["commit", "-q", "--allow-empty", "-m", "a"], ["push", "-q", "origin", "main"],
                 ["checkout", "-qb", "develop"], ["commit", "-q", "--allow-empty", "-m", "b"],
                 ["commit", "-q", "--allow-empty", "-m", "c"], ["push", "-q", "origin", "develop"],
                 ["checkout", "-qb", "feature/y"], ["commit", "-q", "--allow-empty", "-m", "d"],
                 ["push", "-q", "origin", "feature/y"], ["fetch", "-q", "origin"]):
        g(work, *args)
    r0 = {"name": "billing", "branch": "main", "mains": ["develop", "release/9"]}
    check("a feature branch derives from the nearest main branch",
          _repos.base_of(r0, "feature/y", work, str(origin)) == "develop"
          and _repos.base_of(r0, "develop", work, str(origin)) == "main")

    # one embedding per (model, text): a text embedded before never costs another API call
    from brain import llm as _llm
    sent = []
    real_post = _llm._post
    _llm._post = lambda url, key, body, timeout: (sent.append(len(body["input"])), {"data": [
        {"index": i, "embedding": [1.0, float(len(t))]} for i, t in enumerate(body["input"])]})[1]
    cfg.update({"BRAIN_EMBEDDING_MODEL": "m1", "BRAIN_LLM_BASE_URL": "http://gateway/v1"})
    a = _llm.embed(["alpha", "beta"])
    b = _llm.embed(["beta", "gamma"])
    cfg.update({"BRAIN_EMBEDDING_MODEL": "m2"})
    _llm.embed(["alpha"])
    check("vector cache: only new texts reach the API, per model", sent == [2, 1, 1]
          and abs(b[0][1] / b[0][0] - a[1][1] / a[1][0]) < 0.05, (sent, a, b))
    _llm._post = real_post
    cfg.update({"BRAIN_EMBEDDING_MODEL": "", "BRAIN_LLM_BASE_URL": ""})
    try:
        cfg.save_repos([{"name": "a", "path": "x"}], {"s": ["nope"]})
        check("unknown set member rejected", False)
    except ValueError:
        check("unknown set member rejected", True)

    # every document -> Markdown
    import io as _io
    from docx import Document
    from brain.providers import convert
    d = Document(); d.add_heading("Ödeme Kuralları", 1); d.add_paragraph("Banka 503 döndüğünde üç kez denenir.")
    t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Kod"; t.cell(1, 0).text = "503"
    buf = _io.BytesIO(); d.save(buf)
    mdtext = convert.convert("kurallar.docx", buf.getvalue()).text
    check("Word -> Markdown (heading, text, table)", "## Ödeme Kuralları" in mdtext and "| Kod |" in mdtext, mdtext[:160])
    mdtext = convert.convert("p.html", b"<h2>Title</h2><p>Body <code>x()</code></p><script>bad()</script>").text
    check("HTML -> Markdown", "## Title" in mdtext and "`x()`" in mdtext and "bad()" not in mdtext, mdtext)
    check("CSV -> Markdown table", "| a | b |\n|---|---|\n| 1 | 2 |" in convert.convert("t.csv", b"a,b\n1,2\n").text)
    check("code -> fenced block", "```python" in convert.convert("x.py", b"def f(): pass").text)
    try:
        convert.convert("x.bin", b"\x00\x01\x02")
        check("binary rejected", False)
    except ValueError:
        check("binary rejected", True)
    # MarkItDown-style registry: content sniffing, containers and more formats
    mdtext = convert.convert("no-extension", buf.getvalue()).text
    check("Word file without extension recognised by content", "## Ödeme Kuralları" in mdtext, mdtext[:80])
    mdtext = convert.convert("wrong.pdf", buf.getvalue()).text
    check("wrong extension falls back to the content's converter", "## Ödeme Kuralları" in mdtext, mdtext[:80])
    nb = {"nbformat": 4, "metadata": {"kernelspec": {"language": "python"}}, "cells": [
        {"cell_type": "markdown", "source": ["# Analiz"]},
        {"cell_type": "code", "source": ["print(1)"], "outputs": [{"output_type": "stream", "text": ["1\n"]}]}]}
    mdtext = convert.convert("a.ipynb", _json.dumps(nb).encode()).text
    check("notebook -> Markdown", "# Analiz" in mdtext and "```python" in mdtext and "Output:" in mdtext, mdtext)
    eml = (b"From: a@corp\r\nTo: b@corp\r\nSubject: Iade kurali\r\nMIME-Version: 1.0\r\n"
           b"Content-Type: multipart/mixed; boundary=X\r\n\r\n--X\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n"
           b"Refunds need approval.\r\n--X\r\nContent-Type: text/csv\r\nContent-Disposition: attachment; filename=r.csv\r\n\r\n"
           b"a,b\r\n1,2\r\n--X--\r\n")
    mdtext = convert.convert("m.eml", eml).text
    check("e-mail with attachment -> Markdown", "Iade kurali" in mdtext and "Refunds need approval." in mdtext
          and "Attachment: r.csv" in mdtext and "| a | b |" in mdtext, mdtext)
    zbuf = _io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("notes/readme.txt", "hello archive")
        z.writestr("page.html", "<h1>Inside</h1><p>zipped <a href='https://x.corp/a'>link</a></p>")
    mdtext = convert.convert("bundle.zip", zbuf.getvalue()).text
    check("zip -> every file inside", "## notes/readme.txt" in mdtext and "hello archive" in mdtext and "## Inside" in mdtext
          and "[link](https://x.corp/a)" in mdtext, mdtext)
    rss = b"<?xml version='1.0'?><rss><channel><title>News</title><item><title>Yeni kural</title><link>https://n/1</link><description>&lt;p&gt;Text&lt;/p&gt;</description></item></channel></rss>"
    mdtext = convert.convert("feed.xml", rss).text
    check("RSS feed -> sections", "## Yeni kural" in mdtext and "https://n/1" in mdtext, mdtext)
    from brain import repos as _repos
    check("every repository is updated at least every 24 hours", _repos.due(
        ["fresh", "old", "failed", "never"], {"fresh": {"indexed_at": 90_000}, "old": {"indexed_at": 3_000},
                                             "failed": {"indexed_at": 1_000, "failed_at": 2_000}}, 1440, 90_000)
          == ["old", "failed", "never"])
    check("page headers and footers removed", convert._drop_repeated(
        ["ACME Corp - Page 1", "body a", "ACME Corp - Page 2", "body b", "ACME Corp - Page 3", "body c"],
        [{convert._norm_line(f"ACME Corp - Page {i}")} for i in (1, 2, 3)]) == ["body a", "body b", "body c"])

    # requirements writer: formats are only the user's
    from brain import srs
    check("no built-in format", srs.templates() == [], srs.templates())
    head, secs = srs.parse("# {{title}}\nCompany header\n## 1. Purpose\nWhy the system exists.\n## 2. Functional requirements\n"
                           "### 2.1 Per feature\nREQ-? statements\n## 3. Glossary\nTerms.")
    check("user format: headings are the sections", [x["heading"] for x in secs] == ["1. Purpose", "2. Functional requirements",
          "3. Glossary"] and "### 2.1 Per feature" in secs[1]["guidance"] and "Company header" in head, secs)
    head, secs = srs.parse("# Introduction\ntext\n# Requirements\ntext")
    check("converted Word template: level-1 headings are the sections", len(secs) == 2, secs)
    check("requirement IDs numbered per prefix", srs._number("FR-? a FR-? b BR-? REQ-? REQ-XX") == "FR-001 a FR-002 b BR-001 REQ-001 REQ-002")
    check("format without headings rejected", _raises(lambda: srs.save_template("bad", "just text")))
    check("writing without a format explains what to do", _raises(lambda: srs.write("billing")))
    srs.save_template("company", "## Purpose\nWhy.\n## Requirements\nREQ-?")
    check("a saved format is listed and can be deleted", srs.templates() == ["company"]
          and (srs.delete_template("company") or srs.templates() == []))

    # the document engine: structure, concepts, keywords (no model needed)
    from brain.providers import docgraph as dg
    check("Turkish folding", dg.fold("Müşteri İADE Onayı") == "musteri iade onayi" and dg.lower("İade") == "iade")
    parts = dg.passages("# Kurallar\n<!-- page 2 -->\n## İade\nİade onayı gerekir.\n\n| Kod | Anlam |\n|---|---|\n| 503 | Banka |\n"
                        "![Grafik](assets/a.png)\nŞekil 1: İade grafiği", 1500)
    check("passages: headings, pages, tables, pictures", [(p["kind"], p["path"], p["page"]) for p in parts]
          == [("text", "İade", 2), ("table", "İade", 2), ("picture", "İade", 2)] and "Şekil 1" in parts[2]["text"], parts)
    stop = dg._stop()
    found = dg.concepts_in("The Payment Gateway charges VAT via chargeCustomer and \"partial refund\".", stop)
    check("concepts without a model", {"payment gateway", "vat", "chargecustomer", "partial refund"} <= set(found), found)
    check("glossary definitions", dg.definitions_in("text", "- **KDV**: Katma değer vergisi, faturaya eklenir.")
          == {"kdv": ("KDV", "Katma değer vergisi, faturaya eklenir.")})
    eng = dg.Provider()
    eng.start()
    eng.upload("kurallar.md", "# Kurallar\n## İade\nİade onayı için Payment Gateway kullanılır. İade onayı 3 gün sürer.\n".encode())
    eng.upload("refund.md", b"# Refund\nThe Payment Gateway retries. See the Kurallar document.\n")
    r = eng.search("iade onayı", ["iade", "onayı"], 3, 5, ["iade onayı"])
    check("document search: keywords and concepts", r["chunks"] and r["chunks"][0]["doc"] == "kurallar.md"
          and "concept graph" in r["chunks"][0]["matched"] and any(e["name"] == "iade onayı" for e in r["entities"]), r)
    g = eng.graph("*", 50, 1)
    check("document graph: documents, concepts, links", any(n["kind"] == "document" for n in g["nodes"])
          and any(l["type"] == "refers to" for l in g["links"]), g)
    # multi-hop through a shared concept (Personalized PageRank): the second passage shares no word with the question
    eng.upload("atlas.md", b"# Atlas\nProject Atlas is run by the Settlement Desk.\n")
    eng.upload("desk.md", b"# Desk\nThe Settlement Desk approves every chargeback within two days.\n")
    r = eng.search("Project Atlas", ["project", "atlas"], 4, 5, ["project atlas"])
    check("multi-hop: a passage reached only through a shared concept",
          any(c["doc"] == "desk.md" and "concept graph" in c["matched"] and "keywords" not in c["matched"] for c in r["chunks"]),
          [(c["doc"], c["matched"]) for c in r["chunks"]])

    # document understanding without a model: word forms, rules and limits, table rows, aliases by meaning
    families = {"onay": "onay onayı onaylar onayları onayının onaya onaydan", "iade": "iade iadeler iadesi iadenin iadeyi iadeden",
                "müşteri": "müşteri müşteriler müşterinin müşteriye müşterisi", "talep": "talep talebi talepler talebin",
                "hesap": "hesap hesabı hesaplar hesaba", "kart": "kart kartı kartlar karta", "ödeme": "ödeme ödemeler ödemesi ödemeye",
                "yönetmelik": "yönetmelik yönetmeliği yönetmelikler yönetmeliğin", "ona": "ona", "kurum": "kurum kurumu kurumlar"}
    stems = {k: {dg.concept_key(w) for w in v.split()} for k, v in families.items()}
    check("Turkish word forms share one stem; different words stay apart",
          all(len(x) == 1 for x in stems.values()) and len({next(iter(x)) for x in stems.values()}) == len(stems)
          and dg.concept_key("KDV'nin", True, "acronym") == "kdv" and dg.concept_key("Refunds", False) == "refund"
          and dg.concept_key("approved", False) == dg.concept_key("approval", False), stems)
    rules = dg.rules_in("text", "- 5.000 TL üzeri iadeler için müdür onayı zorunludur.\nKartla ödenen faturalar nakit iade "
                                "edilemez.\nRefunds over 1,000 EUR must be approved within 3 business days.\nBu belge süreci anlatır.")
    check("rules: obligation, prohibition, limits; plain sentences skipped",
          [(r["modality"], r["limits"]) for r in rules] == [("obligation", ["5.000 TL üzeri"]), ("prohibition", []),
                                                            ("obligation", ["over 1,000 EUR", "within 3 business days"])], rules)
    rows = dg.rules_in("table", "| Tarife | Limit |\n|---|---|\n| Standart | 5.000 TL |\n| Not | yok |")
    check("table rows with numbers become facts with their column names",
          [r["text"] for r in rows] == ["Tarife: Standart; Limit: 5.000 TL"], rows)
    eng.upload("iade-kurallari.md", ("# İade kuralları\n## Onay\nİade Onayı gerekir. İade Onayının süresi en geç 3 gün "
                                     "içindedir. 5.000 TL üzeri iadeler için müdür onayı zorunludur.\n").encode())
    eng.upload("refunds.md", b"# Refunds\nRefund Approval is recorded by the Payment Gateway.\n")
    keys = {r["key"] for r in eng._q("SELECT key FROM concepts")}
    check("one concept for every word form", "iade onay" in keys and not any(k.startswith("iade onayin") for k in keys), sorted(keys))
    r = eng.search("onaylar", ["onaylar"], 3, 5, [])
    check("keyword search matches other word forms (stemmed column)", any(c["doc"] == "iade-kurallari.md" for c in r["chunks"]),
          [c["doc"] for c in r["chunks"]])
    r = eng.search("müdür onayı zorunlu mu", ["müdür", "onayı", "zorunlu"], 3, 5, [])
    check("search returns the business rules with their limits",
          any(x["modality"] == "obligation" and "5.000 TL üzeri" in x["limits"] for x in r["rules"]), r["rules"])
    import numpy as _np
    real_embed = dg.llm.embed

    def _fake(texts, timeout=None):   # "iade" and "refund" mean the same; everything else is unrelated
        out = []
        for t in texts:
            v, f = _np.zeros(16), dg.fold(t)
            v[0 if ("iade" in f or "refund" in f) else 1 + int(hashlib.md5(f.encode()).hexdigest(), 16) % 15] = 1
            out.append(v.tolist())
        return out
    dg.llm.embed = _fake
    eng._embed_concepts()
    dg.llm.embed = real_embed
    r = eng.search("refund approval", ["refund", "approval"], 5, 8, ["refund approval"])
    ent = {e["name"]: e for e in r["entities"]}
    check("aliases: a Turkish and an English concept that mean the same are joined",
          any(dg.fold(a) == "iade onayi" for a in (ent.get("Refund Approval", {}).get("aliases") or []))
          and any(c["doc"] == "iade-kurallari.md" for c in r["chunks"]), (ent, [c["doc"] for c in r["chunks"]]))

    # pictures: decorative ones skipped, OCR when the vision call fails, the vision model upgrades it, cache by SHA-1
    import io as _io
    from PIL import Image as _Image, ImageDraw as _Draw, ImageFont as _Font
    from brain import ocr as _ocr

    def _png(w, h, text=""):
        img = _Image.new("RGB", (w, h), "white")
        if text:
            _Draw.Draw(img).text((20, 30), text, fill="black", font=_Font.load_default(size=42))
        out = _io.BytesIO()
        img.save(out, "PNG")
        return out.getvalue()
    chart = _png(640, 120, "REFUND LIMIT 5000 EUR")
    eng.upload("dot.png", _png(40, 40))
    eng.upload("chart.png", chart)
    _t.sleep(0.5)
    with eng._completing:   # let the upload's own completion pass finish first
        pass
    kinds = lambda: {r["name"]: r["kind"] for r in eng._q("SELECT d.name, c.kind FROM chunks c JOIN docs d ON d.id = c.doc WHERE c.asset IS NOT NULL")}
    check("pictures: a tiny picture is skipped as decorative", kinds().get("dot.png") == "decor", kinds())
    calls = []
    real_chat = dg.llm.chat

    def _vision_down(*a, **k):
        calls.append(1)
        raise dg.llm.LLMError("timeout after 90s")
    dg.llm.chat = _vision_down
    cfg.update({"BRAIN_VLM_MODEL": "vlm", "BRAIN_LLM_BASE_URL": "http://gateway/v1"})
    eng._describe_pictures()
    text = eng._q("SELECT c.text FROM chunks c JOIN docs d ON d.id = c.doc WHERE d.name = 'chart.png' AND c.asset IS NOT NULL")[0]["text"]
    if _ocr.available():
        check("pictures: OCR reads the text when the vision call fails", kinds()["chart.png"] == "ocr"
              and "REFUND" in text.upper() and calls, (kinds(), text))
    else:
        check("pictures: without OCR a failed vision call leaves the picture waiting", kinds()["chart.png"] == "picture")
    dg.llm.chat = lambda *a, **k: (calls.append(1), {"content": "Bar chart: refund limit 5000 EUR"})[1]
    eng._describe_pictures()
    check("pictures: the vision model's description replaces the OCR text later", kinds()["chart.png"] == "figure"
          and "Bar chart" in eng._q("SELECT c.text FROM chunks c JOIN docs d ON d.id = c.doc WHERE d.name = 'chart.png' AND c.asset IS NOT NULL")[0]["text"]
          and dg.OCR_LABEL not in eng._q("SELECT c.text FROM chunks c JOIN docs d ON d.id = c.doc WHERE d.name = 'chart.png' AND c.asset IS NOT NULL")[0]["text"])
    n = len(calls)
    eng.upload("chart-copy.png", chart)
    _t.sleep(0.5)
    with eng._completing:
        pass
    eng._describe_pictures()
    check("pictures: the same picture in another document reuses its description (no call)",
          kinds()["chart-copy.png"] == "figure" and len(calls) == n, (kinds(), len(calls), n))
    cfg.update({"BRAIN_VLM_MAX_PER_HOUR": "2"})
    dg._vlm_calls.clear()
    check("pictures: hourly vision budget", [dg._vision_budget() for _ in range(3)] == [True, True, False])
    dg.llm.chat = real_chat
    cfg.update({"BRAIN_VLM_MODEL": "", "BRAIN_LLM_BASE_URL": "", "BRAIN_VLM_MAX_PER_HOUR": "60"})
    eng.delete([d["id"] for d in eng.documents() if d["name"].endswith(".png")])

    # document cards: key sentences, dates, versions, numbers, other versions; themes
    base = ("# İade Prosedürü\nDoküman No: PRC-012\nVersiyon 1.0\nYürürlük tarihi: 01.03.2024\n\n## Amaç\nBu prosedür müşteri "
            "iadelerinin nasıl yapılacağını tanımlar. Şube müdürü 5.000 TL üzeri iadeleri onaylamalıdır. İade talepleri en geç 3 iş "
            "günü içinde yanıtlanmalıdır. Kartla yapılan ödemeler nakit olarak iade edilemez. Müşteri bilgileri üçüncü kişilerle "
            "paylaşılmamalıdır. İade kayıtları muhasebe sistemine aynı gün girilir.\n")
    eng.upload("prc-v1.md", base.encode())
    eng.upload("prc-v2.md", base.replace("Versiyon 1.0", "Versiyon 2.0").replace("01.03.2024", "15.09.2025").encode())
    docs = {d["name"]: d for d in eng.documents()}
    check("document card: version, effective date, newer version found",
          docs["prc-v1.md"]["versions"] == ["1.0"] and docs["prc-v1.md"]["effective"] == "2024-03-01"
          and docs["prc-v1.md"]["superseded_by"] == "prc-v2.md" and docs["prc-v2.md"]["superseded_by"] is None, docs.get("prc-v1.md"))
    card = next(c for c in eng.cards().values() if c["name"] == "prc-v2.md")
    check("document card: key sentences and document number", len(card["summary"]) == 4 and card["numbers"] == ["PRC-012"], card)
    r = eng.search("PRC-012 iade prosedürü şube müdürü", ["PRC-012", "prosedürü", "şube", "müdürü"], 6, 5, [])
    order = [c["doc"] for c in r["chunks"]]
    check("the current version ranks above the older one, which says so",
          "prc-v2.md" in order and "prc-v1.md" in order and order.index("prc-v2.md") < order.index("prc-v1.md")
          and all(c.get("note") == "older version; newer: prc-v2.md" for c in r["chunks"] if c["doc"] == "prc-v1.md"),
          [(c["doc"], c.get("note")) for c in r["chunks"]])
    check("themes across documents", any("prc-v1.md" in t["documents"] for t in eng.themes()), eng.themes())

    # document grammar: the unit tree, references, definitions and statements of official and technical documents
    from brain.providers import docgrammar as gr
    reg = ("# Bilgi Güvenliği Yönetmeliği\n\nBİRİNCİ BÖLÜM\nGenel Hükümler\n\nMADDE 1 – Tanımlar\n(1) Bu Yönetmelikte geçen;\n"
           "a) BGYS: Bilgi Güvenliği Yönetim Sistemini,\nb) Sistem sahibi: Bir bilgi sisteminden sorumlu birim yöneticisini,\n"
           "ifade eder.\n\nİKİNCİ BÖLÜM\nErişim\n\nMADDE 2 – Erişim yetkileri\n(1) Sistem sahibi, erişim taleplerini en geç 2 iş "
           "günü içinde onaylamalıdır.\n(2) Ayrıcalıklı hesaplar için iki faktörlü kimlik doğrulama zorunludur.\n\nMADDE 3 – Denetim\n"
           "(1) 2 nci maddenin (2) numaralı fıkrasına aykırılık halinde olay kaydı açılır.\n(2) Kayıtlar 6698 sayılı Kişisel Verilerin "
           "Korunması Kanunu ve ISO/IEC 27001 hükümlerine uygun saklanır.\n\nGEÇİCİ MADDE 1 – (1) Mevcut hesaplar 2 nci maddeye uygun "
           "hale getirilir.\n")
    keys = {u["key"]: u for u in gr.parse_units(reg)}
    check("grammar: chapters, articles, paragraphs, items, provisional articles",
          {"chapter:1", "chapter:2", "art:1/par:1/ite:a", "art:2/par:2", "art:3/par:1", "art:gecici1"} <= set(keys)
          and keys["art:2/par:2"]["path"] == "Bölüm 2 › Madde 2 › (2)", sorted(keys))
    refs = {r["key"] for r in gr.find_refs("2 nci maddenin (2) numaralı fıkrasına göre; birinci fıkranın (a) bendi; Section 3.1(a); "
                                           "REQ-004; 6698 sayılı Kişisel Verilerin Korunması Kanunu; ISO/IEC 27001", {"article": "art:3"})}
    check("grammar: references to articles, paragraphs, items, sections, requirements, laws, standards",
          {"art:2/par:2", "art:3/par:1/ite:a", "sec:3.1/ite:a", "req:REQ-004", "law:6698", "std:ISO/IEC 27001"} <= refs, refs)
    styles = {u["key"] for u in gr.parse_units("# X\n\n**4.3 Acil değişiklik**\nMetin.\n\nARTICLE 2 SERVICE LEVELS\nSection 2.01 The "
                                                "Provider shall keep availability.\n\nA. Purpose\nText.\n")}
    check("grammar: bold numbered paragraphs, contract clauses (Section 2.01), lettered headings",
          {"sec:4.3", "art:2", "sec:2.01", "sec:A"} <= styles, styles)
    check("grammar: what people type finds the unit", gr.unit_keys("madde 5 (2) b")[0] == "art:5/par:2/ite:b"
          and gr.unit_keys("Section 3.1") == ["sec:3.1", "art:3.1"] and gr.unit_keys("REQ-004") == ["req:REQ-004"])
    en = gr.parse_statement("REQ-003: If the PSP is unavailable, the system SHOULD queue the payment, except for refunds.", False)
    tr = gr.parse_statement("(1) Sistem sahibi, erişim taleplerini en geç 2 iş günü içinde onaylamalıdır.", True, "obligation")
    pas = gr.parse_statement("Kayıtlar beş yıl saklanır.", True, "obligation")
    story = gr.parse_statement("As a merchant, I want to see refund status so that I can answer customers.", False)
    check("grammar: statement slots (RFC 2119, condition, exception, Turkish actor, passive, user story)",
          (en["modality"], en["actor"], en["action"], en["condition"], en["exception"]) ==
          ("recommendation", "system", "queue", "If the PSP is unavailable", "except for refunds")
          and (tr["actor"], tr["action"]) == ("Sistem sahibi", "onayla") and pas["actor"] is None and pas["action"] == "saklan"
          and (story["kind"], story["actor"]) == ("story", "merchant"), (en, tr, pas, story))
    acc = gr.gherkin(["Given a refund above 1,000 EUR", "When the manager approves it", "Then the refund is sent within 1 hour"])
    check("grammar: Given / When / Then becomes an acceptance criterion with its condition",
          len(acc) == 1 and acc[0]["condition"] == "a refund above 1,000 EUR; the manager approves it", acc)
    eng.upload("bgy.md", reg.encode())
    out = eng.doc_query("outline", "bgy.md")["text"]
    unit = eng.doc_query("unit", None, unit="Madde 2 (2)")["text"]
    define = eng.doc_query("define", term="BGYS")["text"]
    duties = eng.doc_query("obligations", actor="sistem sahibi")["text"]
    check("doc tool: outline, unit with the units that refer to it, definition, an actor's obligations",
          "Madde 2" in out and "(2)" in out and "referred to by: Bölüm 2 › Madde 3 › (1)" in unit
          and "Bilgi Güvenliği Yönetim Sistemini" in define and "Sistem sahibi" in duties and "onayla" in duties,
          (out, unit, define, duties))

    # decision memory: typed, dated, searchable, replaced by topic, stale when its code disappears
    from brain import memory, providers as prov
    real_docs, prov.docs = prov.docs, (lambda: eng)
    m1 = memory.remember("gotcha", "Refund retries", "Retries stop after three attempts.", "billing", ["src/refund.py"])
    m2 = memory.remember("gotcha", "Refund retries", "Retries stop after five attempts since 2026-10.", "billing", ["src/refund.py"])
    got = memory.recall("refund retries")
    check("memory: replaced by topic, history kept, searchable", not m1["replaced"] and m2["replaced"] and len(got) == 1
          and "five" in got[0]["text"] and got[0]["history"] and any(d["name"] == m2["name"] + ".md" for d in eng.documents()), got)
    check("memory: unknown type rejected", _raises(lambda: memory.remember("idea", "x", "y")))
    (_TMP / "repo" / "src").mkdir(parents=True, exist_ok=True)
    check("memory: stale when its code path is gone", memory.check_stale("billing", _TMP / "repo") == 1
          and memory.recall()[0]["stale"] == ["src/refund.py"])
    memory.forget(m2["name"])
    check("memory: forgotten with its document", memory.recall() == [] and not any(d["name"].startswith("memory-") for d in eng.documents()))
    prov.docs = real_docs

    eng.delete([d["id"] for d in eng.documents()])
    check("documents deleted with everything derived", eng.documents() == [] and eng.labels("") == [])
    eng.stop()

    # search waits at most BRAIN_EMBED_BUDGET for the API and keeps a late embedding for the next search
    import time as _time
    from brain import search as se
    real_embed = se.llm.embed
    se.llm.embed = lambda texts, timeout=None: (_time.sleep(0.6), [[1.0, 0.0] for _ in texts])[1]
    cfg.update({"BRAIN_EMBED_BUDGET": "0.1"})
    t0 = _time.time()
    v1 = se.query_vector("a question on a slow gateway")
    waited = _time.time() - t0
    _time.sleep(0.8)
    v2 = se.query_vector("a question on a slow gateway")
    check("search embedding budget: no long wait, late embedding kept", v1 is None and waited < 0.4 and v2 == [1.0, 0.0],
          (v1, waited, v2))
    se.llm.embed = real_embed
    cfg.update({"BRAIN_EMBED_BUDGET": "3"})

    docs = sorted(p.name for p in (ROOT / "docs" / "decisions").glob("ADR-*.md"))
    check("architecture decision records present", len(docs) >= 27, len(docs))
    whl = [p.name for p in (ROOT / "wheels").glob("*.whl")]
    check("bundled wheels present", len(whl) > 30, len(whl))
    check("wheels are for Python 3.13 only", not [w for w in whl if re.search(r"-cp3(?!13)\d+-cp3", w)])
    ocr_whl = [p.name.lower() for p in (ROOT / "wheels" / "ocr").glob("*.whl")]
    pins = [ln.split("==")[0].strip().lower().replace("-", "_") for ln in
            (ROOT / "requirements-ocr.txt").read_text(encoding="utf-8").splitlines() if "==" in ln]
    check("OCR pack: a wheel for every pin (Windows and Linux where binary)",
          not ocr_whl or all(any(w.replace("-", "_").startswith(p + "_") for w in ocr_whl) for p in pins), pins)
    for f in ("README.md", "INSTALL.html", "setup.ps1", "setup.sh", "brain.cmd", "brain.sh", "requirements.txt",
              "repos.example.json", "bin/SHA256SUMS"):
        check(f"file present: {f}", (ROOT / f).exists())
    for link in re.findall(r"\]\((?!http)([^)#]+)", (ROOT / "README.md").read_text(encoding="utf-8")):
        check(f"README link resolves: {link}", (ROOT / link).exists())
    md_files = sorted((ROOT / "docs").rglob("*.md")) + sorted((ROOT / "integrations").rglob("*.md"))
    broken = [f"{p.relative_to(ROOT)} -> {link}" for p in md_files
              for link in re.findall(r"\]\((?!http)([^)#\s]+)", p.read_text(encoding="utf-8"))
              if not (p.parent / link).exists()]
    check("every link in docs/ resolves", not broken, broken[:5])
    history = re.compile(r"\(added\b|\b[Rr]evised\b|\brevises ADR|\bextends ADR|requested on|second round|reference project")
    stale = [str(p.relative_to(ROOT)) for p in md_files + [ROOT / "README.md", ROOT / "INSTALL.html"]
             if history.search(p.read_text(encoding="utf-8"))]
    check("documentation describes the system as it is (no change-history wording)", not stale, stale)
    print(f"\n{len(FAILS)} failure(s)" + (": " + ", ".join(FAILS) if FAILS else ""))
    sys.exit(len(FAILS))


if __name__ == "__main__":
    main()
