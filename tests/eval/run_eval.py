"""Honest evaluation of the document engine on tests/eval/corpus.

Measures, without any model call unless --with-api is given:
  * indexing: time per document and passages per second, on the corpus and on a copy of it scaled up (--scale N);
  * retrieval: for every question in questions.json, the rank of the first passage that is in an expected document
    and contains the expected phrase: hit@1, hit@5, MRR, per category; the version question checks that the current
    version ranks above the older one;
  * the doc tool: define, obligations, unit, refs, impact answers that must contain an expected text;
  * the grammar against gold.json: units (recall, and precision of numbered units), resolved references, defined
    terms, and statement slots (modality, actor, action).

    python tests/eval/run_eval.py [--scale 30] [--with-api]

--with-api uses the configured corporate embedding model (BRAIN_LLM_BASE_URL, BRAIN_EMBEDDING_MODEL from .env), so
passage and concept vectors take part in search; without it only keywords, concepts and PageRank do.
"""
from __future__ import annotations

import argparse
import json
import os
import statistics
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).parent
ROOT = HERE.parent.parent


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--scale", type=int, default=30, help="copies of the corpus for the speed test (0 = skip)")
    ap.add_argument("--with-api", action="store_true", help="use the configured corporate embedding model")
    a = ap.parse_args()
    tmp = Path(tempfile.mkdtemp())
    os.environ["BRAIN_DATA_DIR"] = str(tmp / "data")
    if not a.with_api:
        os.environ["BRAIN_CONFIG_FILE"] = str(tmp / "eval.env")
        (tmp / "eval.env").write_text("BRAIN_LLM_BASE_URL=\nBRAIN_EMBEDDING_MODEL=\nBRAIN_VLM_MODEL=\n", encoding="utf-8")
    sys.path.insert(0, str(ROOT))
    from brain import query as qa
    from brain.providers import docgraph as dg

    corpus = sorted((HERE / "corpus").iterdir())
    questions = json.loads((HERE / "questions.json").read_text(encoding="utf-8"))
    gold = json.loads((HERE / "gold.json").read_text(encoding="utf-8"))
    eng = dg.Provider()
    eng.start()
    report: dict = {"vectors": "corporate embedding model" if a.with_api else "none (keywords + concepts + PageRank only)"}

    # ---------------------------------------------------------------- indexing
    t0 = time.time()
    passages = 0
    for f in corpus:
        passages += eng.upload(f.name, f.read_bytes())["passages"]
    if a.with_api:
        eng.complete()
    secs = time.time() - t0
    report["indexing"] = {"documents": len(corpus), "passages": passages, "seconds": round(secs, 2),
                          "ms_per_document": round(1000 * secs / len(corpus), 1)}

    # ---------------------------------------------------------------- retrieval
    rows, lat = [], []
    for item in questions["search"]:
        an = qa.analyze(item["q"])
        t = time.time()
        r = eng.search(item["q"], an["keywords"] + an["variants"], 8, 8, an["concepts"])
        lat.append((time.time() - t) * 1000)
        chunks = r["chunks"]
        phrase = dg.fold(item["phrase"])
        rank = next((i + 1 for i, c in enumerate(chunks) if c["doc"] in item["doc"] and phrase in dg.fold(c["content"])), None)
        doc_rank = next((i + 1 for i, c in enumerate(chunks) if c["doc"] in item["doc"]), None)
        ok_version = None
        if item.get("older"):
            order = [c["doc"] for c in chunks]
            cur = order.index(item["doc"][0]) if item["doc"][0] in order else 99
            old = order.index(item["older"]) if item["older"] in order else 99
            ok_version = cur < old and all(c.get("note") for c in chunks if c["doc"] == item["older"])
        rows.append({"id": item["id"], "cat": item["cat"], "q": item["q"], "rank": rank, "doc_rank": doc_rank,
                     "top": chunks[0]["doc"] if chunks else None, "version_ok": ok_version})
    cats: dict[str, list] = {}
    for r in rows:
        cats.setdefault(r["cat"], []).append(r)

    def score(rs):
        n = len(rs)
        return {"n": n, "hit@1": round(sum(1 for r in rs if r["rank"] == 1) / n, 2),
                "hit@5": round(sum(1 for r in rs if r["rank"] and r["rank"] <= 5) / n, 2),
                "mrr": round(sum(1 / r["rank"] for r in rs if r["rank"]) / n, 2),
                "doc@5": round(sum(1 for r in rs if r["doc_rank"] and r["doc_rank"] <= 5) / n, 2)}
    report["retrieval"] = {"all": score(rows), **{c: score(rs) for c, rs in cats.items()},
                           "latency_ms": {"p50": round(statistics.median(lat), 1), "max": round(max(lat), 1)},
                           "misses": [{"id": r["id"], "q": r["q"], "rank": r["rank"], "top": r["top"]} for r in rows
                                      if not r["rank"] or r["rank"] > 5],
                           "version": [{"id": r["id"], "ok": r["version_ok"]} for r in rows if r["version_ok"] is not None]}

    # ---------------------------------------------------------------- doc tool
    tool = []
    for t in questions["tool"]:
        ans = eng.doc_query(t["mode"], t.get("document"), unit=t["q"], term=t["q"], actor=t["q"])["text"]
        tool.append({"id": t["id"], "mode": t["mode"], "ok": dg.fold(t["expect"]) in dg.fold(ans), "answer": ans[:300]})
    report["doc_tool"] = {"passed": sum(t["ok"] for t in tool), "n": len(tool),
                          "failures": [t for t in tool if not t["ok"]]}

    # ---------------------------------------------------------------- grammar
    docs = {r["name"]: r["id"] for r in eng._q("SELECT id, name FROM docs")}
    numbered = ("chapter", "part", "article", "paragraph", "item", "section", "requirement")
    unit_rec, unit_prec = [], []
    for name, keys in gold["units"].items():
        found = {r["key"]: r["kind"] for r in eng._q("SELECT key, kind FROM units WHERE doc = ?", (docs[name],))}
        unit_rec.append((sum(k in found for k in keys), len(keys), [k for k in keys if k not in found]))
        extra = [k for k, kind in found.items() if kind in numbered and not k.startswith("head:") and k not in keys]
        unit_prec.append((len([k for k, kind in found.items() if kind in numbered and not k.startswith("head:")]) - len(extra),
                          len([k for k, kind in found.items() if kind in numbered and not k.startswith("head:")]), extra))
    ref_hits = []
    for name, src, tgt in gold["refs"]:
        hit = eng._q("""SELECT 1 FROM refs r JOIN units s ON s.id = r.unit JOIN units t ON t.id = r.target
            WHERE r.doc = ? AND (s.key = ? OR s.key LIKE ?) AND t.key = ?""", (docs[name], src, src + "/%", tgt))
        ref_hits.append({"ref": f"{name}: {src} -> {tgt}", "ok": bool(hit)})
    def_hits = []
    for name, terms in gold["definitions"].items():
        have = {dg.fold(r[0]) for r in eng._q("SELECT term FROM defs WHERE doc = ?", (docs[name],))}
        have |= {dg.fold(r[0]) for r in eng._q("""SELECT DISTINCT k.name FROM concepts k JOIN mentions m ON m.concept = k.id
            JOIN chunks c ON c.id = m.chunk WHERE c.doc = ? AND k.definition IS NOT NULL""", (docs[name],))}
        for term in terms:
            def_hits.append({"term": f"{name}: {term}", "ok": dg.fold(term) in have})
    st_rows = eng._q("SELECT text, modality, actor, action FROM rules WHERE modality IS NOT NULL")
    slots = {"modality": [0, 0], "actor": [0, 0], "action": [0, 0], "found": [0, 0]}
    st_fail = []
    for g in gold["statements"]:
        slots["found"][1] += 1
        row = next((r for r in st_rows if dg.fold(g["text"]) in dg.fold(r["text"])), None)
        if row is None:
            st_fail.append({"text": g["text"], "problem": "not found as a statement"})
            for k in ("modality", "actor", "action"):
                slots[k][1] += 1
            continue
        slots["found"][0] += 1
        checks = {"modality": row["modality"] == g["modality"],
                  "actor": (not row["actor"]) if g["actor"] is None else bool(row["actor"]) and dg.fold(g["actor"]) in dg.fold(row["actor"]),
                  "action": bool(row["action"]) and dg.fold(g["action"]) in dg.fold(row["action"])}
        for k, ok in checks.items():
            slots[k][0] += ok
            slots[k][1] += 1
        if not all(checks.values()):
            st_fail.append({"text": g["text"], "got": {"modality": row["modality"], "actor": row["actor"], "action": row["action"]},
                            "wrong": [k for k, ok in checks.items() if not ok]})
    report["grammar"] = {
        "units_recall": f"{sum(a for a, _, _ in unit_rec)}/{sum(b for _, b, _ in unit_rec)}",
        "units_missing": [m for _, _, ms in unit_rec for m in ms],
        "numbered_units_precision": f"{sum(a for a, _, _ in unit_prec)}/{sum(b for _, b, _ in unit_prec)}",
        "units_extra": [e for _, _, es in unit_prec for e in es],
        "references_resolved": f"{sum(r['ok'] for r in ref_hits)}/{len(ref_hits)}",
        "references_missing": [r["ref"] for r in ref_hits if not r["ok"]],
        "definitions_found": f"{sum(d['ok'] for d in def_hits)}/{len(def_hits)}",
        "definitions_missing": [d["term"] for d in def_hits if not d["ok"]],
        "statements": {k: f"{v[0]}/{v[1]}" for k, v in slots.items()},
        "statement_errors": st_fail}

    # ---------------------------------------------------------------- robustness: styles the corpus does not use
    report["robustness"] = robustness(eng)

    # ---------------------------------------------------------------- scale
    if a.scale:
        t0 = time.time()
        n = p = 0
        for k in range(a.scale):
            for f in corpus:
                p += eng.upload(f"copy{k:03d}-{f.name}", f.read_bytes())["passages"]
                n += 1
        secs = time.time() - t0
        lat = []
        for item in questions["search"]:
            an = qa.analyze(item["q"])
            t = time.time()
            eng.search(item["q"], an["keywords"] + an["variants"], 8, 8, an["concepts"])
            lat.append((time.time() - t) * 1000)
        report["scale"] = {"documents": n + len(corpus), "added": n, "passages_added": p, "seconds": round(secs, 1),
                           "documents_per_second": round(n / secs, 1), "passages_per_second": round(p / secs, 1),
                           "search_ms_p50": round(statistics.median(lat), 1), "search_ms_max": round(max(lat), 1)}
    eng.stop()
    (HERE / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 0


def robustness(eng) -> dict:
    """Unit recall on three styles the corpus does not use: a Word file whose headings are bold normal paragraphs,
    an English contract with "ARTICLE 1" and "Section 1.01" clauses, and lettered headings ("A. Purpose")."""
    import io
    import docx
    sys.path.insert(0, str(HERE))
    import make_corpus
    d = docx.Document()
    for line in make_corpus.DOCS["degisiklik-yonetimi-proseduru-v2.docx"].splitlines():
        if line.strip() and not line.startswith("|"):
            d.add_paragraph().add_run(line.lstrip("# ")).bold = line.startswith("#")
    buf = io.BytesIO()
    d.save(buf)
    contract = ("MASTER SERVICES AGREEMENT\n\nARTICLE 1 DEFINITIONS\nSection 1.01 \"Services\" means the hosting services.\n"
                "Section 1.02 \"Downtime\" means any period in which the Services are unavailable.\n\nARTICLE 2 SERVICE LEVELS\n"
                "Section 2.01 The Provider shall keep monthly availability at or above 99.9%.\nSection 2.02 If Downtime exceeds the "
                "limit in Section 2.01, the Customer may claim a credit of 5% of the monthly fee.\n\nARTICLE 3 TERMINATION\n"
                "Section 3.01 Either party may terminate this Agreement with 90 days written notice.\n")
    lettered = ("IT Asset Policy\n\nA. Purpose\nThis policy defines how laptops are issued and returned.\n\nB. Responsibilities\n"
                "Employees must return laptops on their last day.\n")
    cases = {"plain-paragraphs.docx": (buf.getvalue(), ["sec:1", "sec:2", "sec:3", "sec:4", "sec:4.1", "sec:4.2", "sec:4.3", "sec:5",
                                                        "sec:5.1", "sec:5.2", "sec:6"]),
             "contract.md": (contract.encode(), ["art:1", "art:2", "art:3", "sec:1.01", "sec:1.02", "sec:2.01", "sec:2.02", "sec:3.01"]),
             "lettered.md": (lettered.encode(), ["sec:A", "sec:B"])}
    out = {}
    for name, (data, want) in cases.items():
        eng.upload(name, data)
        have = {r[0] for r in eng._q("SELECT u.key FROM units u JOIN docs d ON d.id = u.doc WHERE d.name = ?", (name,))}
        out[name] = f"{sum(k in have for k in want)}/{len(want)}"
    return out


if __name__ == "__main__":
    sys.exit(main())
