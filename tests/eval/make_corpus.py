"""Builds the evaluation corpus: twelve official documents of a fictional technology company ("Kuzey Teknoloji"),
in Turkish and English, as Word, PDF, Excel and Markdown files, written in the styles such companies use (regulation
with articles, numbered procedures in two versions, an SRS with requirement IDs, a policy, a runbook, standards, a
retention schedule, meeting notes, a glossary). Run once; the files go to tests/eval/corpus/.

    python tests/eval/make_corpus.py
"""
from __future__ import annotations

import re
from pathlib import Path

OUT = Path(__file__).parent / "corpus"

DOCS: dict[str, str] = {}

DOCS["bilgi-guvenligi-yonetmeligi.pdf"] = """Kuzey Teknoloji A.Ş. Bilgi Güvenliği Yönetmeliği
Doküman No: YNT-001
Versiyon 3.0
Yürürlük tarihi: 01.02.2026

BİRİNCİ BÖLÜM
Amaç, Kapsam, Dayanak ve Tanımlar

MADDE 1 – Amaç
(1) Bu Yönetmeliğin amacı, Şirketin bilgi varlıklarının gizlilik, bütünlük ve erişilebilirliğinin korunmasına ilişkin usul ve esasları belirlemektir.

MADDE 2 – Kapsam
(1) Bu Yönetmelik, Şirket çalışanlarını, stajyerleri ve Şirket sistemlerine erişen tedarikçi personelini kapsar.

MADDE 3 – Dayanak
(1) Bu Yönetmelik, 6698 sayılı Kişisel Verilerin Korunması Kanunu ve ISO/IEC 27001 standardı esas alınarak hazırlanmıştır.

MADDE 4 – Tanımlar
(1) Bu Yönetmelikte geçen;
a) Bilgi varlığı: Şirket için değeri olan her türlü veri, yazılım, donanım ve hizmeti,
b) BGYS: Bilgi Güvenliği Yönetim Sistemini,
c) Sistem sahibi: Bir bilgi sisteminin işletiminden sorumlu birim yöneticisini,
ç) Ayrıcalıklı hesap: Sistem yapılandırmasını değiştirme yetkisine sahip kullanıcı hesabını,
ifade eder.

İKİNCİ BÖLÜM
Erişim Yönetimi

MADDE 5 – Erişim talepleri
(1) Sistem sahibi, erişim taleplerini en geç 2 iş günü içinde onaylamalı veya reddetmelidir.
(2) Ayrıcalıklı hesaplar için iki faktörlü kimlik doğrulama zorunludur. Ancak acil durum hesapları, Bilgi Güvenliği Birimi onayı ile bu fıkranın kapsamı dışında tutulabilir.
(3) Erişim hakları, personelin görevi sona erdiğinde aynı gün kaldırılır.

MADDE 6 – Parola kuralları
(1) Parolalar en az 12 karakter uzunluğunda olmalıdır.
(2) Parolalar 90 günde bir değiştirilmelidir.
(3) Aynı parola son 5 parola içinde tekrar kullanılamaz.

ÜÇÜNCÜ BÖLÜM
Denetim ve Yaptırımlar

MADDE 7 – Erişim gözden geçirmesi
(1) Bilgi Güvenliği Birimi tarafından ayrıcalıklı hesaplar üç ayda bir gözden geçirilmelidir.
(2) 5 inci maddenin (2) numaralı fıkrasına aykırılık tespit edilmesi halinde olay kaydı açılır ve Olay Müdahale Planı uygulanır.

MADDE 8 – Kayıtların saklanması
(1) Erişim kayıtları, Veri Saklama Politikasında belirtilen sürelerle ve 6698 sayılı Kanun hükümlerine uygun olarak saklanır.

MADDE 9 – Yaptırımlar
(1) Bu Yönetmelik hükümlerine aykırı davranan personel hakkında İnsan Kaynakları Disiplin Prosedürü uygulanır.

GEÇİCİ MADDE 1 – (1) Mevcut ayrıcalıklı hesaplar 31.03.2026 tarihine kadar 5 inci maddeye uygun hale getirilir.

MADDE 10 – Yürürlük
(1) Bu Yönetmelik 01.02.2026 tarihinde yürürlüğe girer.
(2) Bilgi Güvenliği Yönetmeliği Versiyon 2.0 yürürlükten kaldırılmıştır.
"""

CHANGE = """Değişiklik Yönetimi Prosedürü
Doküman No: PRS-014
Versiyon {version}
Yürürlük tarihi: {date}

# 1. Amaç
Bu prosedür, üretim ortamındaki sistemlerde yapılacak değişikliklerin planlanması, onaylanması ve uygulanması için izlenecek adımları tanımlar.

# 2. Kapsam
Bu prosedür, Kuzey Teknoloji bünyesindeki tüm üretim sistemlerini, veri tabanlarını ve ağ cihazlarını kapsar.

# 3. Roller ve Sorumluluklar
| Rol | Sorumluluk |
|---|---|
| Değişiklik Sahibi | Değişiklik talebini oluşturur ve uygular |
| Değişiklik Danışma Kurulu | Normal değişiklikleri değerlendirir ve onaylar |
| Değişiklik Yöneticisi | Prosedürün uygulanmasını takip eder |

# 4. Değişiklik Türleri
## 4.1 Standart değişiklik
Standart değişiklikler önceden onaylanmış, düşük riskli ve tekrarlanan işlemlerdir. Standart değişiklikler için ayrıca onay gerekmez.
## 4.2 Normal değişiklik
Normal değişiklikler Değişiklik Danışma Kurulu tarafından haftalık toplantıda onaylanmalıdır. Değişiklik talebi en geç uygulamadan 3 iş günü önce açılmalıdır.
## 4.3 Acil değişiklik
Acil değişiklikler, Değişiklik Yöneticisi tarafından en geç {hours} saat içinde onaylanmalıdır. Acil değişiklikler sonradan Değişiklik Danışma Kuruluna raporlanır.{extra}

# 5. Uygulama
## 5.1 Geri dönüş planı
Her normal ve acil değişiklik için geri dönüş planı hazırlanmalıdır.
## 5.2 Uygulama penceresi
Normal değişiklikler yalnızca hafta içi 22:00 ile 02:00 arasında uygulanabilir. 4.3 numaralı bölümde tanımlanan acil değişiklikler bu kısıtlamanın dışındadır.

# 6. İlgili Dokümanlar
Bilgi Güvenliği Yönetmeliği (YNT-001), Olay Müdahale Planı (PLN-003).
"""
DOCS["degisiklik-yonetimi-proseduru-v1.docx"] = CHANGE.format(version="1.0", date="01.03.2025", hours="4", extra="")
DOCS["degisiklik-yonetimi-proseduru-v2.docx"] = CHANGE.format(
    version="2.0", date="01.06.2026", hours="2",
    extra="\nVeri tabanı şema değişiklikleri acil değişiklik olarak uygulanamaz.")

DOCS["olay-mudahale-plani.md"] = """# Olay Müdahale Planı
Doküman No: PLN-003
Versiyon 1.2
Yürürlük tarihi: 15.04.2026

## 1. Amaç
Bu plan, bilgi güvenliği olaylarının tespit edilmesi, sınıflandırılması ve çözülmesi için uygulanacak adımları belirler.

## 2. Öncelik seviyeleri
| Öncelik | Tanım | İlk müdahale süresi | Çözüm hedefi |
|---|---|---|---|
| P1 | Müşteri hizmetinin tamamen durması veya veri sızıntısı | 15 dakika | 4 saat |
| P2 | Önemli bir özelliğin çalışmaması | 1 saat | 1 gün |
| P3 | Sınırlı etki, geçici çözüm mevcut | 4 saat | 5 iş günü |
| P4 | Bilgi talebi veya kozmetik sorun | 1 iş günü | 15 iş günü |

## 3. Roller
Olay Yöneticisi, olayın koordinasyonundan ve paydaşların bilgilendirilmesinden sorumludur. Nöbetçi mühendis, ilk müdahaleyi yapar ve olayı sınıflandırır.

## 4. Müdahale adımları
1) Nöbetçi mühendis olayı 15 dakika içinde olay yönetim sistemine kaydetmelidir.
2) P1 ve P2 olaylarda Olay Yöneticisi atanmalıdır.
3) Kişisel veri sızıntısı şüphesi varsa Hukuk Birimi ve Veri Koruma Sorumlusu derhal bilgilendirilmelidir.
4) Kişisel veri ihlalleri, öğrenildiği andan itibaren en geç 72 saat içinde Kişisel Verileri Koruma Kuruluna bildirilmelidir.
5) Olay kapatıldıktan sonra 5 iş günü içinde olay sonrası değerlendirme raporu yazılmalıdır.

## 5. İlgili dokümanlar
Bilgi Güvenliği Yönetmeliği (YNT-001), Değişiklik Yönetimi Prosedürü (PRS-014).
"""

DOCS["payments-srs.md"] = """# Payments Platform – Software Requirements Specification
Document No: SRS-PAY-001
Version 2.1
Effective date: 15.09.2026

## 1. Introduction
This specification defines the requirements of the Kuzey Payments Platform. The key words MUST, SHOULD and MAY are to be interpreted as described in RFC 2119.

## 2. Definitions
| Term | Meaning |
|---|---|
| PSP | Payment service provider that processes card payments |
| Chargeback | A payment reversal initiated by the card holder's bank |

"Settlement Window" means the daily period in which captured payments are transferred to merchants.

## 3. Functional requirements
### 3.1 Payments
REQ-001: The Payment Service MUST retry a failed PSP call at most 3 times within 60 seconds.
REQ-002: The Payment Service SHALL NOT store full card numbers.
REQ-003: If the PSP is unavailable, the system SHOULD queue the payment and notify the on-call engineer.

### 3.2 Refunds
REQ-004: Refunds above 1,000 EUR MUST be approved by a finance manager. This requirement depends on REQ-001 and Section 3.1.
REQ-005: The Refund Service MAY issue partial refunds, except for gift card payments.

As a merchant, I want to see refund status in the dashboard so that I can answer customer questions.

Given a refund above 1,000 EUR
When the finance manager approves it
Then the refund is sent to the PSP within 1 hour

## 4. Non-functional requirements
NFR-001: The API MUST respond within 300 ms for 95% of requests.
NFR-002: Card data MUST be processed according to PCI DSS and ISO/IEC 27001.
NFR-003: The platform SHOULD be available 99.95% of the time per month.

## 5. Change history
This version supersedes SRS-PAY-001 version 2.0. Section 3.2 is amended to add REQ-005.
"""

DOCS["access-control-policy.docx"] = """Access Control Policy
Document No: POL-007
Version 1.4
Effective date: 01.07.2026

# 1. Purpose
This policy sets the rules for granting, reviewing and removing access to Kuzey Teknoloji systems.

# 2. Scope
This policy applies to all employees, contractors and service accounts.

# 3. Principles
## 3.1 Least privilege
Users shall be granted only the access they need to perform their job.
## 3.2 Segregation of duties
The same person shall not both request and approve an access change.

# 4. Account management
## 4.1 New accounts
Line managers must approve new accounts before they are created.
## 4.2 Privileged accounts
Privileged accounts must use multi-factor authentication. Privileged sessions should be recorded.
## 4.3 Leavers
Access of leavers must be removed on their last working day, except for e-mail forwarding approved by HR.

# 5. Reviews
The Security Team shall review all privileged accounts every quarter. Findings that violate Section 4.2 must be fixed within 10 business days.

# 6. References
ISO/IEC 27001 Annex A, Information Security Regulation (YNT-001).
"""

DOCS["api-tasarim-standartlari.md"] = """# API Tasarım Standartları
Doküman No: STD-021
Versiyon 1.0
Yürürlük tarihi: 10.01.2026

## 1. Genel kurallar
1) Tüm servisler REST prensiplerine uygun tasarlanmalıdır.
2) API sürümü URL içinde belirtilmelidir, örneğin /v1/payments.
3) Kaynak isimleri çoğul ve küçük harfle yazılmalıdır.

## 2. Hata yönetimi
| HTTP kodu | Kullanım |
|---|---|
| 400 | Geçersiz istek |
| 401 | Kimlik doğrulama başarısız |
| 404 | Kaynak bulunamadı |
| 429 | Hız sınırı aşıldı |

Hata yanıtları code, message ve traceId alanlarını içermelidir.

## 3. Güvenlik
Tüm uç noktalar OAuth 2.0 ile korunmalıdır. Kişisel veri içeren yanıtlar loglanmamalıdır.

## 4. Performans
Liste uç noktaları sayfalama desteklemelidir ve tek istekte en fazla 100 kayıt döndürmelidir.
"""

DOCS["veri-saklama-politikasi.xlsx"] = "TABLE"   # built separately as an Excel workbook

DOCS["kod-inceleme-standardi.md"] = """# Kod İnceleme Standardı
Doküman No: STD-030
Versiyon 2.0
Yürürlük tarihi: 05.05.2026

## 1. Amaç
Bu standart, ana dala birleştirilecek kod değişikliklerinin incelenme kurallarını belirler.

## 2. Kurallar
Her birleştirme isteği en az iki geliştirici tarafından onaylanmalıdır. Güvenlikle ilgili değişikliklerde onaylardan biri Güvenlik Ekibinden olmalıdır.
Birleştirme isteğini açan kişi kendi isteğini onaylayamaz.
Birim testleri başarısız olan birleştirme istekleri birleştirilemez.
İncelemeler en geç 1 iş günü içinde tamamlanmalıdır.

## 3. Kapsam dışı
Yalnızca dokümantasyon içeren değişikliklerde tek onay yeterlidir.
"""

DOCS["release-runbook.md"] = """# Release Runbook
Document No: RB-002
Version 1.1

## 1. Overview
This runbook describes how the Platform Team ships a release to production every Tuesday.

## 2. Steps
1) The release manager creates the release branch on Monday.
2) QA must sign off the regression suite before 12:00 on Tuesday.
3) The release manager deploys to production between 14:00 and 16:00.
4) If error rates rise above 2% within 30 minutes, the on-call engineer must roll back the release.

## 3. Responsibilities
| Activity | Responsible | Accountable |
|---|---|---|
| Release branch | Release manager | Engineering manager |
| Regression sign-off | QA lead | Engineering manager |
| Rollback | On-call engineer | Release manager |
"""

DOCS["toplanti-notlari-2026-09-12.md"] = """# Mimari Kurul Toplantı Notları – 12 Eylül 2026
Katılımcılar: Ayşe K., Mehmet T., Can D., Zeynep A.

Gündemde ödeme servisinin yeni veri tabanına taşınması vardı. Mehmet, mevcut veri tabanının yoğun saatlerde yavaşladığını ve bağlantı havuzunun sık sık dolduğunu anlattı. Can, PostgreSQL 16'ya geçişin bu sorunu çözeceğini düşünüyor.

Karar: Ödeme servisi Ekim sonuna kadar PostgreSQL 16'ya taşınacak. Taşıma işi Değişiklik Yönetimi Prosedürüne göre normal değişiklik olarak planlanacak.

Ayrıca Zeynep, mobil uygulamada oturum süresinin 30 dakikadan 15 dakikaya düşürülmesini önerdi; konu bir sonraki toplantıya bırakıldı.
"""

DOCS["sozluk.md"] = """# Kurumsal Sözlük / Corporate Glossary

| Terim | Açıklama |
|---|---|
| DDK | Değişiklik Danışma Kurulu; normal değişiklikleri onaylayan kurul (Change Advisory Board) |
| KVKK | 6698 sayılı Kişisel Verilerin Korunması Kanunu |
| SLA | Hizmet seviyesi anlaşması (Service Level Agreement) |
| RTO | Bir hizmetin kesintiden sonra en geç ne kadar sürede geri gelmesi gerektiği (Recovery Time Objective) |
| MFA | Çok faktörlü kimlik doğrulama (multi-factor authentication) |
"""

RETENTION = [("Veri Türü", "Saklama Süresi", "Dayanak", "Sorumlu Birim"),
             ("Erişim kayıtları (log)", "2 yıl", "Bilgi Güvenliği Yönetmeliği Madde 8", "Bilgi Güvenliği Birimi"),
             ("Müşteri sözleşmeleri", "10 yıl", "Türk Ticaret Kanunu", "Hukuk Birimi"),
             ("Ödeme işlem kayıtları", "10 yıl", "6493 sayılı Kanun", "Finans"),
             ("CV ve iş başvuruları", "6 ay", "KVKK", "İnsan Kaynakları"),
             ("Kamera kayıtları", "30 gün", "KVKK", "İdari İşler")]


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, text in DOCS.items():
        path = OUT / name
        if name.endswith(".md"):
            path.write_text(text, encoding="utf-8")
        elif name.endswith(".docx"):
            _docx(text, path)
        elif name.endswith(".pdf"):
            _pdf(text, path)
        elif name.endswith(".xlsx"):
            _xlsx(path)
    print(f"{len(DOCS)} documents in {OUT}")


def _docx(text: str, path: Path) -> None:
    import docx
    d = docx.Document()
    lines = text.strip().splitlines()
    d.add_heading(lines[0], level=0)
    table: list[list[str]] = []
    for line in lines[1:] + [""]:
        if line.startswith("|"):
            if not re.match(r"\|\s*-", line):
                table.append([c.strip() for c in line.strip("|").split("|")])
            continue
        if table:
            t = d.add_table(rows=len(table), cols=len(table[0]))
            t.style = "Table Grid"
            for r, row in enumerate(table):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
            table = []
        if m := re.match(r"^(#+)\s+(.*)$", line):
            d.add_heading(m.group(2), level=len(m.group(1)))
        elif line.strip():
            d.add_paragraph(line)
    d.save(path)


def _pdf(text: str, path: Path) -> None:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.add_font("Arial", "", r"C:\Windows\Fonts\arial.ttf")
    pdf.add_font("Arial", "B", r"C:\Windows\Fonts\arialbd.ttf")
    for line in text.strip().splitlines():
        bold = bool(re.match(r"^(MADDE|GEÇİCİ MADDE|[A-ZÇĞİÖŞÜ]+ BÖLÜM)", line)) and len(line) < 40
        pdf.set_font("Arial", "B" if bold else "", 11)
        pdf.multi_cell(0, 6, line or " ", new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def _xlsx(path: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Saklama Süreleri"
    for row in RETENTION:
        ws.append(row)
    wb.save(path)


if __name__ == "__main__":
    main()
